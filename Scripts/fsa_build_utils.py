#!/usr/bin/env python3
"""
Shared helpers for the FSA Title IV volume-report panel pipeline.

The design mirrors the newer IPEDSDB_Panel repository:
- repo-local code and tests
- durable data and QA artifacts outside git under FSA_ROOT
- stage scripts that mostly coordinate shared logic from this module
"""
from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Sequence
from urllib.parse import unquote, urljoin, urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_FSA_ROOT = Path("/Users/markjaysonfarol13/Projects/FSAVolumeReports_Paneling")
REPO_ROOT = Path(__file__).resolve().parents[1]
TITLE_IV_PAGE_URL = "https://studentaid.gov/data-center/student/title-iv"
TITLE_IV_JSON_URL = f"{TITLE_IV_PAGE_URL}.json"
USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 14_4) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36"
)

VALID_DOWNLOAD_SUFFIXES = {".xls", ".xlsx", ".csv"}
DEFINITION_SHEET_HINTS = ("definition", "field definitions")
SUMMARY_SHEET_HINTS = ("award year summary", "summary", "ytd", "cumulative")
DESCRIPTOR_COLUMN_TOKENS = {
    "ope_id",
    "opeid",
    "school",
    "state",
    "zip_code",
    "school_type",
}

COMPONENT_FAMILIES = {
    "grants": {
        "display": "Grant Volume",
        "download_dir": "Grant_Volume",
        "cross_sections_dir": "grants",
        "panel_dir": "grants",
        "prefix": "grant__",
    },
    "campus_based": {
        "display": "Campus-Based Volume",
        "download_dir": "Campus_Based_Volume",
        "cross_sections_dir": "campus_based",
        "panel_dir": "campus_based",
        "prefix": "campus__",
    },
    "direct_loans": {
        "display": "Loan Volume / Direct",
        "download_dir": "Loan_Volume/Direct",
        "cross_sections_dir": "loans_direct",
        "panel_dir": "loans",
        "prefix": "loan_direct__",
    },
    "ffel": {
        "display": "Loan Volume / FFEL",
        "download_dir": "Loan_Volume/FFEL",
        "cross_sections_dir": "loans_ffel",
        "panel_dir": "loans",
        "prefix": "loan_ffel__",
    },
}

AWARD_YEAR_PATTERNS: Sequence[re.Pattern[str]] = [
    re.compile(r"AY(?P<start>\d{4})[-_](?P<end>\d{2,4})", flags=re.IGNORECASE),
    re.compile(r"AY(?P<start>\d{2})[-_](?P<end>\d{2})", flags=re.IGNORECASE),
    re.compile(r"(?P<start>\d{4})[-_](?P<end>\d{2,4})"),
    re.compile(r"(?P<start>\d{2})[-_](?P<end>\d{2})"),
]

COMPONENT_ORDER = ["grants", "campus_based", "direct_loans", "ffel"]

EXPECTED_SELECTED_SCOPE = {
    "grants": {
        "annual_years": tuple(range(1999, 2006)),
        "q4_years": tuple(range(2006, 2025)),
    },
    "campus_based": {
        "annual_years": tuple(range(2001, 2024)),
        "q4_years": (),
    },
    "direct_loans": {
        "annual_years": tuple(range(1999, 2006)),
        "q4_years": tuple(range(2006, 2025)),
    },
    "ffel": {
        "annual_years": tuple(range(1999, 2006)),
        "q4_years": tuple(range(2006, 2010)),
    },
}

INVENTORY_FIELDS = [
    "family",
    "family_display",
    "link_text",
    "url",
    "filename",
    "award_year",
    "award_year_start",
    "award_year_end",
    "quarter",
    "period_type",
    "selected_for_panel",
    "selection_reason",
]


@dataclass(frozen=True)
class DataRootLayout:
    root: Path
    raw_title_iv_reports: Path
    cross_sections: Path
    dictionary: Path
    panels: Path
    checks: Path
    build: Path


@dataclass
class WorkbookParseResult:
    workbook_path: Path
    selected_sheet: str
    sheet_names: list[str]
    header_row_index: int | None
    flattened_headers: list[str]
    frame: pd.DataFrame
    warnings: list[str]


def repo_root() -> Path:
    return REPO_ROOT


def data_root() -> Path:
    return Path(os.environ.get("FSA_ROOT", str(DEFAULT_FSA_ROOT))).expanduser()


def data_layout(root: str | Path | None = None) -> DataRootLayout:
    base = Path(root).expanduser() if root is not None else data_root()
    return DataRootLayout(
        root=base,
        raw_title_iv_reports=base / "Raw_Title_IV_Reports",
        cross_sections=base / "Cross_sections",
        dictionary=base / "Dictionary",
        panels=base / "Panels",
        checks=base / "Checks",
        build=base / "build",
    )


def ensure_data_layout(root: str | Path | None = None) -> DataRootLayout:
    layout = data_layout(root)
    for path in (
        layout.root,
        layout.raw_title_iv_reports,
        layout.cross_sections,
        layout.dictionary,
        layout.panels,
        layout.checks,
        layout.build,
        layout.checks / "download_qc",
        layout.checks / "source_qc",
        layout.checks / "dictionary_qc",
        layout.checks / "panel_qc",
        layout.checks / "acceptance_qc",
        layout.checks / "logs",
    ):
        path.mkdir(parents=True, exist_ok=True)
    for family in COMPONENT_FAMILIES:
        component_download_root(layout, family).mkdir(parents=True, exist_ok=True)
        component_cross_sections_dir(layout, family).mkdir(parents=True, exist_ok=True)
        component_panel_dir(layout, family).mkdir(parents=True, exist_ok=True)
    (layout.panels / "final").mkdir(parents=True, exist_ok=True)
    return layout


def component_download_root(layout: DataRootLayout, family: str) -> Path:
    return layout.raw_title_iv_reports / COMPONENT_FAMILIES[family]["download_dir"]


def component_year_dir(layout: DataRootLayout, family: str, start_year: int) -> Path:
    return component_download_root(layout, family) / str(int(start_year))


def component_cross_sections_dir(layout: DataRootLayout, family: str) -> Path:
    return layout.cross_sections / COMPONENT_FAMILIES[family]["cross_sections_dir"]


def component_panel_dir(layout: DataRootLayout, family: str) -> Path:
    if family == "loans":
        return layout.panels / "loans"
    return layout.panels / COMPONENT_FAMILIES[family]["panel_dir"]


def write_rows(path: Path, rows: list[dict], fieldnames: Sequence[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(fieldnames))
        writer.writeheader()
        if rows:
            writer.writerows(rows)


def safe_text(value: object) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:  # noqa: BLE001
        pass
    return str(value)


def compute_file_metadata(path: Path) -> tuple[int, str]:
    if not path.exists():
        return 0, ""
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            if chunk:
                digest.update(chunk)
    return path.stat().st_size, digest.hexdigest()


def normalize_token(value: object) -> str:
    text = safe_text(value).strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    text = re.sub(r"_+", "_", text)
    return text.strip("_")


def normalize_year_component(value: object, reference: int | None = None) -> int | None:
    digits = re.sub(r"\D", "", safe_text(value))
    if not digits:
        return None
    if len(digits) == 4:
        return int(digits)
    if len(digits) == 2:
        yy = int(digits)
        if reference is not None:
            century = reference // 100
            candidate = (century * 100) + yy
            if candidate < reference:
                candidate += 100
            return candidate
        return 1900 + yy if yy >= 90 else 2000 + yy
    return None


def normalize_award_year_string(label: object) -> str | None:
    text = safe_text(label).strip()
    if not text:
        return None
    for pattern in AWARD_YEAR_PATTERNS:
        match = pattern.search(text)
        if not match:
            continue
        start_year = normalize_year_component(match.group("start"))
        end_year = normalize_year_component(match.group("end"), reference=start_year)
        if start_year is None or end_year is None:
            continue
        return f"{start_year:04d}-{end_year:04d}"
    return None


def explode_award_year_bounds(label: object) -> tuple[int | None, int | None]:
    normalized = normalize_award_year_string(label)
    if normalized is None:
        return None, None
    start_txt, end_txt = normalized.split("-", 1)
    return int(start_txt), int(end_txt)


def infer_quarter(text: object) -> int | None:
    match = re.search(r"(?:^|[^0-9A-Z])Q(?P<q>[1-4])(?:[^0-9A-Z]|$)", safe_text(text), flags=re.IGNORECASE)
    if not match:
        return None
    return int(match.group("q"))


def standardize_opeid8(raw: object) -> str | None:
    digits = re.sub(r"\D", "", safe_text(raw))
    if not digits:
        return None
    if set(digits) == {"0"}:
        return None
    if len(digits) >= 8:
        standardized = digits[:8]
        return None if set(standardized) == {"0"} else standardized
    if len(digits) == 7:
        standardized = digits.zfill(8)
        return None if set(standardized) == {"0"} else standardized
    standardized = f"{digits.zfill(6)}00"
    return None if set(standardized) == {"0"} else standardized


def derive_opeid6(opeid8: object) -> str | None:
    standardized = standardize_opeid8(opeid8)
    if standardized is None:
        return None
    return standardized[:6]


def sanitize_href(href: str) -> str:
    cleaned = safe_text(href).strip().strip('"').strip("'")
    return cleaned


def basename_from_url(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme == "file":
        return Path(unquote(parsed.path)).name
    return Path(parsed.path).name


def classify_report_family(url: str, link_text: str) -> str | None:
    filename = basename_from_url(url)
    suffix = Path(filename).suffix.lower()
    if suffix not in VALID_DOWNLOAD_SUFFIXES:
        return None
    haystack = normalize_token(f"{filename} {link_text}")
    if any(token in haystack for token in ("campus_based", "campusbased", "cbdashboard", "cbdata")):
        return "campus_based"
    if any(token in haystack for token in ("ffel_awardyr", "ffel_award", "fl_dashboard")):
        return "ffel"
    if any(token in haystack for token in ("dl_dashboard", "dl_awardyr", "dl_award")):
        return "direct_loans"
    if any(token in haystack for token in ("grants_ay", "pell", "q4", "q3", "q2", "q1")) and (
        "grant" in haystack or "pell" in haystack or re.search(r"q[1-4]\d{4}ay", filename, flags=re.IGNORECASE)
    ):
        return "grants"
    return None


def append_title_iv_entry(
    rows: list[dict],
    seen: set[tuple[str, str]],
    *,
    url: str,
    text: str,
) -> None:
    family = classify_report_family(url, text)
    if family is None:
        return
    filename = basename_from_url(url)
    award_year = normalize_award_year_string(f"{text} {filename}")
    if award_year is None:
        return
    quarter = infer_quarter(f"{text} {filename}")
    key = (family, url)
    if key in seen:
        return
    seen.add(key)
    start_year, end_year = explode_award_year_bounds(award_year)
    period_type = "quarterly" if quarter is not None else "annual"
    rows.append(
        {
            "family": family,
            "family_display": COMPONENT_FAMILIES[family]["display"],
            "link_text": text,
            "url": url,
            "filename": filename,
            "award_year": award_year,
            "award_year_start": start_year,
            "award_year_end": end_year,
            "quarter": quarter or "",
            "period_type": period_type,
        }
    )


def sort_inventory_rows(rows: list[dict]) -> list[dict]:
    return sorted(
        rows,
        key=lambda row: (
            COMPONENT_ORDER.index(row["family"]),
            int(row["award_year_start"]),
            0 if row["quarter"] == "" else int(row["quarter"]),
            row["filename"],
        ),
    )


def parse_title_iv_entries(html: str, base_url: str = TITLE_IV_PAGE_URL) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    rows: list[dict] = []
    seen: set[tuple[str, str]] = set()
    for anchor in soup.find_all("a", href=True):
        href = sanitize_href(anchor["href"])
        if not href:
            continue
        url = urljoin(base_url, href)
        text = anchor.get_text(" ", strip=True) or basename_from_url(url)
        append_title_iv_entry(rows, seen, url=url, text=text)
    return sort_inventory_rows(rows)


def parse_title_iv_json(payload: object, base_url: str = TITLE_IV_PAGE_URL) -> list[dict]:
    rows: list[dict] = []
    seen: set[tuple[str, str]] = set()

    def walk(node: object) -> None:
        if isinstance(node, dict):
            if node.get("widgetType") == "anchor" and node.get("href"):
                url = urljoin(base_url, sanitize_href(str(node.get("href", ""))))
                text = str(node.get("value", "")).strip() or basename_from_url(url)
                append_title_iv_entry(rows, seen, url=url, text=text)
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    walk(payload)
    return sort_inventory_rows(rows)


def annotate_panel_selection(entries: Iterable[dict]) -> list[dict]:
    annotated: list[dict] = []
    for row in entries:
        entry = dict(row)
        family = entry["family"]
        start_year = int(entry["award_year_start"])
        quarter = int(entry["quarter"]) if str(entry["quarter"]).strip() else None
        selected = False
        reason = ""
        if family == "campus_based":
            if quarter is not None:
                selected = False
                reason = "quarterly_not_expected_for_campus_based"
            elif start_year <= 2023:
                selected = True
                reason = "annual_campus_based_selected"
            else:
                selected = False
                reason = "award_year_not_complete"
        elif family == "ffel":
            if quarter is None and start_year <= 2005:
                selected = True
                reason = "annual_summary_selected"
            elif quarter == 4 and 2006 <= start_year <= 2009:
                selected = True
                reason = "q4_cumulative_selected"
            elif quarter in {1, 2, 3}:
                selected = False
                reason = "non_q4_quarter_excluded"
            else:
                selected = False
                reason = "award_year_not_complete"
        else:
            if quarter is None and start_year <= 2005:
                selected = True
                reason = "annual_summary_selected"
            elif quarter == 4 and start_year <= 2024:
                selected = True
                reason = "q4_cumulative_selected"
            elif quarter in {1, 2, 3}:
                selected = False
                reason = "non_q4_quarter_excluded"
            else:
                selected = False
                reason = "award_year_not_complete"
        entry["selected_for_panel"] = selected
        entry["selection_reason"] = reason
        annotated.append(entry)
    return annotated


def load_title_iv_page_html(page_html: str | Path | None = None, timeout: int = 120) -> tuple[str, str]:
    if page_html is not None:
        path = Path(page_html).expanduser()
        return path.read_text(encoding="utf-8"), path.parent.as_uri() + "/"
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    response = session.get(TITLE_IV_PAGE_URL, timeout=timeout)
    response.raise_for_status()
    return response.text, TITLE_IV_PAGE_URL


def load_title_iv_page_json(timeout: int = 120) -> object:
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    response = session.get(TITLE_IV_JSON_URL, timeout=timeout)
    response.raise_for_status()
    return response.json()


def local_source_path_from_url(url: str) -> Path | None:
    parsed = urlparse(url)
    if parsed.scheme == "file":
        return Path(unquote(parsed.path))
    if parsed.scheme == "" and Path(url).exists():
        return Path(url)
    return None


def download_file(session: requests.Session, url: str, destination: Path, timeout: int) -> None:
    source_path = local_source_path_from_url(url)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if source_path is not None:
        shutil.copyfile(source_path, destination)
        return
    with session.get(url, stream=True, timeout=timeout) as response:
        response.raise_for_status()
        with destination.open("wb") as fh:
            for chunk in response.iter_content(chunk_size=65536):
                if chunk:
                    fh.write(chunk)


def manifest_rows(layout: DataRootLayout) -> list[dict]:
    rows: list[dict] = []
    for manifest_path in sorted(layout.raw_title_iv_reports.glob("**/manifest.csv")):
        frame = pd.read_csv(manifest_path, dtype=str).fillna("")
        rows.extend(frame.to_dict("records"))
    return rows


def release_inventory_path(layout: DataRootLayout) -> Path:
    return layout.checks / "download_qc" / "release_inventory.csv"


def selected_panel_files_path(layout: DataRootLayout) -> Path:
    return layout.checks / "download_qc" / "selected_panel_files.csv"


def preflight_release_inventory_path(layout: DataRootLayout) -> Path:
    return layout.checks / "download_qc" / "preflight_release_inventory.csv"


def preflight_selected_panel_files_path(layout: DataRootLayout) -> Path:
    return layout.checks / "download_qc" / "preflight_selected_panel_files.csv"


def preflight_summary_path(layout: DataRootLayout) -> Path:
    return layout.checks / "download_qc" / "preflight_inventory_summary.csv"


def preflight_validation_path(layout: DataRootLayout) -> Path:
    return layout.checks / "download_qc" / "preflight_validation.csv"


def as_bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    return safe_text(value).strip().lower() in {"1", "true", "t", "yes", "y"}


def discover_title_iv_inventory(*, page_html: str | Path | None = None, timeout: int = 120) -> list[dict]:
    if page_html is not None:
        path = Path(page_html).expanduser()
        if path.suffix.lower() == ".json":
            payload = json.loads(path.read_text(encoding="utf-8"))
            return annotate_panel_selection(parse_title_iv_json(payload, base_url=TITLE_IV_PAGE_URL))
        html, base_url = load_title_iv_page_html(page_html=page_html, timeout=timeout)
        return annotate_panel_selection(parse_title_iv_entries(html, base_url=base_url))
    try:
        payload = load_title_iv_page_json(timeout=timeout)
        rows = parse_title_iv_json(payload, base_url=TITLE_IV_PAGE_URL)
        if rows:
            return annotate_panel_selection(rows)
    except Exception:  # noqa: BLE001
        pass
    html, base_url = load_title_iv_page_html(page_html=None, timeout=timeout)
    return annotate_panel_selection(parse_title_iv_entries(html, base_url=base_url))


def selected_inventory_rows(entries: Iterable[dict]) -> list[dict]:
    return [dict(row) for row in entries if as_bool(row.get("selected_for_panel"))]


def expected_selected_start_years(family: str) -> list[int]:
    scope = EXPECTED_SELECTED_SCOPE[family]
    return sorted({*scope["annual_years"], *scope["q4_years"]})


def expected_selected_file_count(family: str) -> int:
    scope = EXPECTED_SELECTED_SCOPE[family]
    return len(scope["annual_years"]) + len(scope["q4_years"])


def summarize_inventory(entries: Iterable[dict]) -> pd.DataFrame:
    rows: list[dict] = []
    for family in COMPONENT_ORDER:
        family_entries = [row for row in entries if row["family"] == family]
        selected = selected_inventory_rows(family_entries)
        discovered_years = sorted({int(row["award_year_start"]) for row in family_entries})
        selected_years = sorted(int(row["award_year_start"]) for row in selected)
        selected_quarters = sorted(
            {int(row["quarter"]) for row in selected if str(row.get("quarter", "")).strip()}
        )
        rows.append(
            {
                "family": family,
                "family_display": COMPONENT_FAMILIES[family]["display"],
                "discovered_files": int(len(family_entries)),
                "selected_files": int(len(selected)),
                "expected_selected_files": int(expected_selected_file_count(family)),
                "discovered_start_year_min": "" if not discovered_years else int(discovered_years[0]),
                "discovered_start_year_max": "" if not discovered_years else int(discovered_years[-1]),
                "selected_start_year_min": "" if not selected_years else int(min(selected_years)),
                "selected_start_year_max": "" if not selected_years else int(max(selected_years)),
                "selected_quarters": json.dumps(selected_quarters),
                "selected_years_json": json.dumps(selected_years),
                "expected_selected_years_json": json.dumps(expected_selected_start_years(family)),
            }
        )
    return pd.DataFrame(rows)


def validate_inventory_scope(entries: Iterable[dict]) -> pd.DataFrame:
    rows: list[dict] = []
    for family in COMPONENT_ORDER:
        family_entries = [row for row in entries if row["family"] == family]
        selected = selected_inventory_rows(family_entries)
        actual_selected_years = sorted(int(row["award_year_start"]) for row in selected)
        expected_years = expected_selected_start_years(family)
        duplicate_years = sorted(
            {
                year
                for year in actual_selected_years
                if actual_selected_years.count(year) > 1
            }
        )
        invalid_selected_quarters = sorted(
            {
                f"{row['award_year']}::Q{row['quarter']}"
                for row in selected
                if str(row.get("quarter", "")).strip() and int(row["quarter"]) != 4
            }
        )
        out_of_scope_selected_years = sorted(set(actual_selected_years) - set(expected_years))
        rows.extend(
            [
                {
                    "check": "family_present_in_discovery",
                    "component_family": family,
                    "passed": bool(family_entries),
                    "details": f"discovered_files={len(family_entries)}",
                },
                {
                    "check": "selected_award_years_match_expected_scope",
                    "component_family": family,
                    "passed": actual_selected_years == expected_years,
                    "details": f"expected={expected_years}; actual={actual_selected_years}",
                },
                {
                    "check": "selected_file_count_matches_expected_scope",
                    "component_family": family,
                    "passed": len(selected) == expected_selected_file_count(family),
                    "details": f"expected={expected_selected_file_count(family)}; actual={len(selected)}",
                },
                {
                    "check": "selected_rows_unique_by_award_year",
                    "component_family": family,
                    "passed": not duplicate_years,
                    "details": f"duplicate_start_years={duplicate_years}",
                },
                {
                    "check": "selected_rows_use_only_annual_or_q4_files",
                    "component_family": family,
                    "passed": not invalid_selected_quarters,
                    "details": f"invalid_selected_quarters={invalid_selected_quarters}",
                },
                {
                    "check": "selected_rows_do_not_include_out_of_scope_years",
                    "component_family": family,
                    "passed": not out_of_scope_selected_years,
                    "details": f"out_of_scope_selected_start_years={out_of_scope_selected_years}",
                },
            ]
        )
    return pd.DataFrame(rows)


def write_preflight_inventory(layout: DataRootLayout, entries: list[dict]) -> tuple[pd.DataFrame, pd.DataFrame]:
    write_rows(preflight_release_inventory_path(layout), entries, INVENTORY_FIELDS)
    selected = selected_inventory_rows(entries)
    write_rows(preflight_selected_panel_files_path(layout), selected, INVENTORY_FIELDS)
    summary = summarize_inventory(entries)
    summary.to_csv(preflight_summary_path(layout), index=False)
    validation = validate_inventory_scope(entries)
    validation.to_csv(preflight_validation_path(layout), index=False)
    return summary, validation


def preflight_title_iv_inventory(
    root: str | Path | None = None,
    *,
    page_html: str | Path | None = None,
    timeout: int = 120,
    strict: bool = True,
) -> tuple[list[dict], pd.DataFrame, pd.DataFrame]:
    layout = ensure_data_layout(root)
    entries = discover_title_iv_inventory(page_html=page_html, timeout=timeout)
    summary, validation = write_preflight_inventory(layout, entries)
    failed = validation.loc[~validation["passed"].astype(bool)]
    if strict and not failed.empty:
        raise SystemExit(
            "Preflight inventory validation failed. Inspect "
            f"{preflight_validation_path(layout)} and {preflight_selected_panel_files_path(layout)}."
        )
    return entries, summary, validation


def download_title_iv_reports(
    root: str | Path | None = None,
    *,
    page_html: str | Path | None = None,
    skip_existing: bool = True,
    timeout: int = 120,
    strict_source_checks: bool = True,
    verify_only: bool = False,
) -> list[dict]:
    layout = ensure_data_layout(root)
    entries, _, _ = preflight_title_iv_inventory(
        layout.root,
        page_html=page_html,
        timeout=timeout,
        strict=strict_source_checks,
    )
    if verify_only:
        return entries
    write_rows(release_inventory_path(layout), entries, INVENTORY_FIELDS)

    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    failures: list[dict] = []
    manifest_map: dict[tuple[str, int], list[dict]] = {}
    for entry in entries:
        family = entry["family"]
        start_year = int(entry["award_year_start"])
        year_dir = component_year_dir(layout, family, start_year)
        download_dir = year_dir / "downloads"
        download_dir.mkdir(parents=True, exist_ok=True)
        destination = download_dir / entry["filename"]
        status = "existing"
        try:
            if not (skip_existing and destination.exists()):
                download_file(session, entry["url"], destination, timeout)
                status = "downloaded"
        except Exception as exc:  # noqa: BLE001
            status = "failed"
            failures.append(
                {
                    "family": family,
                    "award_year_start": start_year,
                    "filename": entry["filename"],
                    "error": str(exc),
                }
            )
        size_bytes, sha256 = compute_file_metadata(destination)
        manifest_row = {
            **entry,
            "local_path": str(destination),
            "download_status": status,
            "filesize_bytes": size_bytes,
            "sha256": sha256,
            "downloaded_at": datetime.now(timezone.utc).isoformat(),
        }
        manifest_map.setdefault((family, start_year), []).append(manifest_row)
    manifest_fields = [
        *INVENTORY_FIELDS,
        "local_path",
        "download_status",
        "filesize_bytes",
        "sha256",
        "downloaded_at",
    ]
    for (family, start_year), rows in manifest_map.items():
        write_rows(component_year_dir(layout, family, start_year) / "manifest.csv", rows, manifest_fields)
    write_rows(
        layout.checks / "download_qc" / "download_failures.csv",
        failures,
        ["family", "award_year_start", "filename", "error"],
    )
    selected_rows = [row for row in manifest_rows(layout) if str(row.get("selected_for_panel", "")).lower() == "true"]
    write_rows(selected_panel_files_path(layout), selected_rows, manifest_fields)
    if failures:
        raise SystemExit(f"Download failures encountered for {len(failures)} files.")
    return entries


def load_selected_panel_entries(layout: DataRootLayout, family: str | None = None) -> list[dict]:
    path = selected_panel_files_path(layout)
    if not path.exists():
        raise SystemExit(f"Missing selected panel manifest: {path}")
    frame = pd.read_csv(path, dtype=str).fillna("")
    if family is not None:
        frame = frame[frame["family"] == family].copy()
    return frame.to_dict("records")


def choose_selected_sheet(
    sheet_names: Sequence[str],
    family: str,
    period_type: str,
) -> str:
    if not sheet_names:
        raise ValueError("Workbook has no sheets.")
    candidates = [name for name in sheet_names if normalize_token(name) not in {"", "sheet1"}]
    if not candidates:
        candidates = list(sheet_names)
    filtered = [
        name
        for name in candidates
        if not any(hint in normalize_token(name) for hint in ("definitions", "field_definitions", "definition"))
    ]
    candidates = filtered or candidates
    if family == "campus_based":
        return candidates[0]
    if period_type == "quarterly":
        for name in candidates:
            normalized = normalize_token(name)
            if any(hint in normalized for hint in ("award_year_summary", "summary", "ytd", "cumulative")):
                return name
    return candidates[0]


def read_workbook_raw(path: Path, sheet_name: str | None = None) -> tuple[list[str], pd.DataFrame, str]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        frame = pd.read_csv(path, header=None, dtype=str, keep_default_na=False)
        return ["__csv__"], frame, "__csv__"
    with pd.ExcelFile(path) as workbook:
        selected_sheet = choose_selected_sheet(
            workbook.sheet_names,
            family="campus_based" if sheet_name == "__campus_first__" else "grants",
            period_type="annual",
        )
        if sheet_name and sheet_name not in {"__campus_first__", ""}:
            selected_sheet = sheet_name
        raw = pd.read_excel(path, sheet_name=selected_sheet, header=None, dtype=str)
        return workbook.sheet_names, raw.fillna(""), selected_sheet


def find_header_row(frame: pd.DataFrame) -> int | None:
    for idx in range(min(len(frame), 20)):
        tokens = {normalize_token(value) for value in frame.iloc[idx].tolist()}
        if "ope_id" in tokens or "opeid" in tokens:
            return idx
    return None


def propagated_group_labels(values: Sequence[object]) -> list[str]:
    propagated: list[str] = []
    current = ""
    for value in values:
        normalized = safe_text(value).strip()
        if normalized:
            current = normalized
        propagated.append(current)
    return propagated


def flattened_headers_from_rows(
    group_row: Sequence[object] | None,
    header_row: Sequence[object],
) -> list[str]:
    group_labels = propagated_group_labels(group_row or [""] * len(header_row))
    headers: list[str] = []
    seen: dict[str, int] = {}
    for group_label, header_label in zip(group_labels, header_row):
        group_label = safe_text(group_label).strip()
        header_label = safe_text(header_label).strip()
        header_norm = normalize_token(header_label)
        base = ""
        if header_norm in DESCRIPTOR_COLUMN_TOKENS or group_label == "":
            base = header_label or group_label or "unnamed"
        elif header_label == "":
            base = group_label
        elif normalize_token(group_label) == normalize_token(header_label):
            base = group_label
        else:
            base = f"{group_label} | {header_label}"
        base = re.sub(r"\s+", " ", base).strip()
        if not base:
            base = "unnamed"
        count = seen.get(base, 0)
        seen[base] = count + 1
        if count:
            base = f"{base}_{count}"
        headers.append(base)
    return headers


def read_html_table_fallback(path: Path) -> pd.DataFrame | None:
    try:
        text = path.read_text(encoding="utf-8", errors="ignore")
    except Exception:  # noqa: BLE001
        return None
    if "<table" not in text.lower():
        return None
    soup = BeautifulSoup(text, "html.parser")
    table = soup.find("table")
    if table is None:
        return None
    rows: list[list[str]] = []
    for tr in table.find_all("tr"):
        cells = tr.find_all(["th", "td"])
        if not cells:
            continue
        rows.append([cell.get_text(" ", strip=True) for cell in cells])
    if not rows:
        return None
    width = max(len(row) for row in rows)
    padded_rows = [row + [""] * (width - len(row)) for row in rows]
    return pd.DataFrame(padded_rows).fillna("")


def text_file_hint(path: Path, limit: int = 256) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="ignore")[:limit].strip()
    except Exception:  # noqa: BLE001
        return ""


def parse_selected_sheet(
    workbook_path: Path,
    family: str,
    period_type: str,
) -> WorkbookParseResult:
    warnings: list[str] = []
    suffix = workbook_path.suffix.lower()
    if suffix == ".csv":
        raw = pd.read_csv(workbook_path, header=None, dtype=str, keep_default_na=False).fillna("")
        selected_sheet = "__csv__"
        sheet_names = [selected_sheet]
    else:
        try:
            with pd.ExcelFile(workbook_path) as workbook:
                sheet_names = workbook.sheet_names
                try:
                    selected_sheet = choose_selected_sheet(sheet_names, family=family, period_type=period_type)
                except ValueError as exc:
                    warnings.append(str(exc))
                    return WorkbookParseResult(
                        workbook_path=workbook_path,
                        selected_sheet="",
                        sheet_names=sheet_names,
                        header_row_index=None,
                        flattened_headers=[],
                        frame=pd.DataFrame(),
                        warnings=warnings,
                    )
                raw = pd.read_excel(workbook_path, sheet_name=selected_sheet, header=None, dtype=str).fillna("")
        except Exception as exc:  # noqa: BLE001
            html_frame = read_html_table_fallback(workbook_path)
            if html_frame is not None:
                warnings.append(f"excel_open_failed:{type(exc).__name__}:html_table_fallback")
                raw = html_frame
                selected_sheet = "__html_table__"
                sheet_names = [selected_sheet]
            else:
                hint = text_file_hint(workbook_path)
                if hint:
                    warnings.append(f"text_file_hint:{hint[:120]}")
                warnings.append(f"excel_open_failed:{type(exc).__name__}")
                return WorkbookParseResult(
                    workbook_path=workbook_path,
                    selected_sheet="",
                    sheet_names=[],
                    header_row_index=None,
                    flattened_headers=[],
                    frame=pd.DataFrame(),
                    warnings=warnings,
                )
    header_row_index = find_header_row(raw)
    if header_row_index is None:
        warnings.append("header_row_not_found")
        return WorkbookParseResult(
            workbook_path=workbook_path,
            selected_sheet=selected_sheet,
            sheet_names=sheet_names,
            header_row_index=None,
            flattened_headers=[],
            frame=pd.DataFrame(),
            warnings=warnings,
        )
    group_row = raw.iloc[header_row_index - 1].tolist() if header_row_index > 0 else None
    header_row = raw.iloc[header_row_index].tolist()
    flattened_headers = flattened_headers_from_rows(group_row, header_row)
    frame = raw.iloc[header_row_index + 1 :].copy()
    frame.columns = flattened_headers
    frame = frame.replace({"": pd.NA}).dropna(how="all").reset_index(drop=True)
    return WorkbookParseResult(
        workbook_path=workbook_path,
        selected_sheet=selected_sheet,
        sheet_names=sheet_names,
        header_row_index=header_row_index,
        flattened_headers=flattened_headers,
        frame=frame,
        warnings=warnings,
    )


def profile_workbooks(root: str | Path | None = None) -> pd.DataFrame:
    layout = ensure_data_layout(root)
    rows: list[dict] = []
    for manifest_row in manifest_rows(layout):
        local_path = Path(manifest_row["local_path"])
        result = parse_selected_sheet(local_path, manifest_row["family"], manifest_row["period_type"])
        opeids = pd.Series(dtype="string")
        non00_rows = 0
        if not result.frame.empty:
            opeid_col = next(
                (col for col in result.frame.columns if normalize_token(col) in {"ope_id", "opeid"}),
                None,
            )
            if opeid_col is not None:
                opeids = result.frame[opeid_col].apply(standardize_opeid8).astype("string")
                opeids = opeids.dropna()
                non00_rows = int((opeids.str[-2:] != "00").sum())
        signature = hashlib.sha256(
            "|".join(normalize_token(col) for col in result.flattened_headers).encode("utf-8")
        ).hexdigest()
        rows.append(
            {
                "family": manifest_row["family"],
                "award_year": manifest_row["award_year"],
                "award_year_start": manifest_row["award_year_start"],
                "filename": manifest_row["filename"],
                "local_path": manifest_row["local_path"],
                "selected_for_panel": manifest_row["selected_for_panel"],
                "selected_sheet": result.selected_sheet,
                "sheet_names_json": json.dumps(result.sheet_names),
                "header_row_index": "" if result.header_row_index is None else result.header_row_index,
                "flattened_header_signature": signature,
                "flattened_headers_json": json.dumps(result.flattened_headers),
                "row_count": int(len(result.frame)),
                "opeid_count": int(len(opeids)),
                "non00_opeid_count": non00_rows,
                "parse_warnings_json": json.dumps(result.warnings),
            }
        )
    profile_path = layout.checks / "source_qc" / "workbook_profiles.csv"
    write_rows(
        profile_path,
        rows,
        [
            "family",
            "award_year",
            "award_year_start",
            "filename",
            "local_path",
            "selected_for_panel",
            "selected_sheet",
            "sheet_names_json",
            "header_row_index",
            "flattened_header_signature",
            "flattened_headers_json",
            "row_count",
            "opeid_count",
            "non00_opeid_count",
            "parse_warnings_json",
        ],
    )
    return pd.DataFrame(rows)


def normalize_metric_header(header: str) -> str:
    text = safe_text(header).strip()
    text = text.replace("#", " number ")
    text = text.replace("$", " amount ")
    normalized = normalize_token(text)
    normalized = normalized.replace("recipents", "recipients")
    normalized = normalized.replace("recipent", "recipient")
    normalized = normalized.replace("sum_of_", "")
    normalized = normalized.replace("ytd_", "")
    normalized = normalized.replace("federal_", "federal_")
    return normalized


def metric_kind_from_header(normalized_header: str) -> str | None:
    if "recipient" in normalized_header:
        return "recipients"
    if "federal_award" in normalized_header:
        return "federal_award"
    if "number_of_loans_originated" in normalized_header or "num_of_loans_originated" in normalized_header:
        return "loans_originated_n"
    if "of_loans_originated" in normalized_header or "amount_of_loans_originated" in normalized_header:
        return "loans_originated_amt"
    if "number_of_disbursements" in normalized_header or "num_of_disbursements" in normalized_header:
        return "disbursements_n"
    if "of_disbursements" in normalized_header or "amount_of_disbursements" in normalized_header or "disbursement" in normalized_header or "disbursements" in normalized_header or "total_disbursed" in normalized_header:
        if "number_of_" in normalized_header or "num_of_" in normalized_header:
            return "disbursements_n"
        return "disbursements_amt" if "loan" in normalized_header else "disbursements"
    return None


def program_slug_from_header(family: str, normalized_header: str) -> str | None:
    if family == "grants":
        if "federal_pell_grant" in normalized_header:
            return "pell"
        if "academic_competitiveness" in normalized_header:
            return "acg"
        if "national_smart" in normalized_header or "smart_program" in normalized_header:
            return "smart"
        if "teach_program" in normalized_header or normalized_header.startswith("teach_"):
            return "teach"
        if "iraq" in normalized_header or "afghanistan" in normalized_header or "iasg" in normalized_header:
            return "iasg"
    elif family == "campus_based":
        if "fseog" in normalized_header or "federal_supplemental_educational_opportunity_grants" in normalized_header:
            return "fseog"
        if "fws" in normalized_header or "federal_work_study" in normalized_header:
            return "fws"
        if "perkins" in normalized_header:
            return "perkins"
    elif family in {"direct_loans", "ffel"}:
        if "unsubsidized_undergraduate" in normalized_header:
            return "unsubsidized_undergraduate"
        if "unsubsidized_graduate" in normalized_header:
            return "unsubsidized_graduate"
        if "unsubsidized" in normalized_header:
            return "unsubsidized"
        if "subsidized_undergraduate" in normalized_header:
            return "subsidized_undergraduate"
        if "subsidized_graduate" in normalized_header:
            return "subsidized_graduate"
        if "subsidized" in normalized_header:
            return "subsidized"
        if "parent_plus" in normalized_header:
            return "parent_plus"
        if "grad_plus" in normalized_header or "graduate_plus" in normalized_header:
            return "grad_plus"
        if re.search(r"(^|_)plus($|_)", normalized_header):
            return "plus"
    return None


def canonicalize_component_header(family: str, original_header: str) -> str | None:
    normalized = normalize_metric_header(original_header)
    if normalized in {"ope_id", "opeid"}:
        return "opeid8"
    if normalized == "school":
        return "school"
    if normalized == "state":
        return "state"
    if normalized == "zip_code":
        return "zip_code"
    if normalized == "school_type":
        return "school_type"
    metric_kind = metric_kind_from_header(normalized)
    program_slug = program_slug_from_header(family, normalized)
    if metric_kind is None or program_slug is None:
        return None
    return f"{program_slug}_{metric_kind}"


def series_has_nonempty_values(series: pd.Series) -> bool:
    cleaned = series.astype("string").str.strip().replace({"": pd.NA, "nan": pd.NA, "None": pd.NA})
    return bool(cleaned.notna().any())


def build_dictionary(root: str | Path | None = None) -> pd.DataFrame:
    layout = ensure_data_layout(root)
    entries = load_selected_panel_entries(layout)
    rows: list[dict] = []
    unmapped_rows: list[dict] = []
    for entry in entries:
        path = Path(entry["local_path"])
        result = parse_selected_sheet(path, entry["family"], entry["period_type"])
        for header in result.flattened_headers:
            series = result.frame[header] if header in result.frame.columns else pd.Series(dtype="string")
            canonical_column = canonicalize_component_header(entry["family"], header)
            normalized_header = normalize_metric_header(header)
            actionable = series_has_nonempty_values(series) and normalized_header != "unnamed"
            mapping_status = "mapped" if canonical_column else ("unmapped_actionable" if actionable else "nonactionable")
            row = {
                "component_family": entry["family"],
                "source_filename": entry["filename"],
                "sheet_name": result.selected_sheet,
                "original_header": header,
                "normalized_header": normalized_header,
                "canonical_column": canonical_column or "",
                "mapping_status": mapping_status,
                "first_year": int(entry["award_year_start"]),
                "last_year": int(entry["award_year_end"]),
            }
            rows.append(row)
            if mapping_status == "unmapped_actionable":
                unmapped_rows.append(row)
    detail = pd.DataFrame(rows)
    if detail.empty:
        raise SystemExit("Dictionary build found no selected panel files.")
    aggregated = (
        detail.groupby(["component_family", "original_header", "normalized_header", "canonical_column", "mapping_status"], dropna=False)
        .agg(
            source_filename=("source_filename", lambda s: ";".join(sorted(set(str(v) for v in s if str(v).strip())))),
            sheet_name=("sheet_name", lambda s: ";".join(sorted(set(str(v) for v in s if str(v).strip())))),
            first_year=("first_year", "min"),
            last_year=("last_year", "max"),
        )
        .reset_index()
        .sort_values(["component_family", "canonical_column", "original_header"])
        .reset_index(drop=True)
    )
    dict_path = layout.dictionary / "fsa_volume_dictionary.parquet"
    aggregated.to_parquet(dict_path, index=False)
    aggregated.to_csv(layout.dictionary / "fsa_volume_dictionary.csv", index=False)
    unmapped = pd.DataFrame(unmapped_rows)
    if unmapped.empty:
        unmapped = pd.DataFrame(columns=aggregated.columns)
    unmapped.to_csv(layout.checks / "dictionary_qc" / "unmapped_actionable_headers.csv", index=False)
    summary_rows = [
        {"metric": "dictionary_rows", "value": int(len(aggregated))},
        {"metric": "mapped_rows", "value": int((aggregated["mapping_status"] == "mapped").sum())},
        {"metric": "unmapped_actionable_rows", "value": int((aggregated["mapping_status"] == "unmapped_actionable").sum())},
        {"metric": "nonactionable_rows", "value": int((aggregated["mapping_status"] == "nonactionable").sum())},
    ]
    write_rows(layout.checks / "dictionary_qc" / "dictionary_qaqc_summary.csv", summary_rows, ["metric", "value"])
    return aggregated


def load_dictionary_map(layout: DataRootLayout) -> dict[tuple[str, str], str]:
    path = layout.dictionary / "fsa_volume_dictionary.parquet"
    if not path.exists():
        raise SystemExit(f"Missing dictionary artifact: {path}")
    frame = pd.read_parquet(path)
    mapping: dict[tuple[str, str], str] = {}
    for row in frame.to_dict("records"):
        if row.get("canonical_column"):
            mapping[(row["component_family"], row["normalized_header"])] = row["canonical_column"]
    return mapping


def sanitize_numeric_series(series: pd.Series) -> pd.Series:
    cleaned = series.astype("string").str.strip()
    cleaned = cleaned.str.replace(",", "", regex=False)
    cleaned = cleaned.str.replace("$", "", regex=False)
    cleaned = cleaned.str.replace("(", "-", regex=False).str.replace(")", "", regex=False)
    cleaned = cleaned.replace({"": pd.NA, "nan": pd.NA, "None": pd.NA})
    return pd.to_numeric(cleaned, errors="coerce")


def format_integer_series(series: pd.Series) -> pd.Series:
    numeric = sanitize_numeric_series(series)
    return numeric.round().astype("Int64")


def format_string_series(series: pd.Series, formatter: str | None = None) -> pd.Series:
    cleaned = series.astype("string").str.strip()
    cleaned = cleaned.replace({"": pd.NA, "nan": pd.NA, "None": pd.NA})
    if formatter == "upper":
        cleaned = cleaned.str.upper()
    elif formatter == "title":
        cleaned = cleaned.str.title()
    return cleaned


def format_zip_series(series: pd.Series) -> pd.Series:
    def _format(value: object) -> object:
        if pd.isna(value):
            return pd.NA
        digits = re.sub(r"\D", "", str(value))
        if not digits:
            return pd.NA
        if len(digits) >= 9:
            digits = digits[:9]
            return f"{digits[:5]}-{digits[5:]}"
        return digits[:5].zfill(5)

    return series.apply(_format).astype("string")


def combine_duplicate_series(existing: pd.Series, incoming: pd.Series, *, label: str) -> pd.Series:
    existing = existing.reset_index(drop=True)
    incoming = incoming.reset_index(drop=True)
    overlap = existing.notna() & incoming.notna()
    if overlap.any():
        left = existing[overlap].astype("string")
        right = incoming[overlap].astype("string")
        if not left.equals(right):
            raise ValueError(f"Conflicting duplicate values encountered for {label}.")
    return existing.combine_first(incoming)


def apply_component_mapping(
    family: str,
    frame: pd.DataFrame,
    dictionary_map: dict[tuple[str, str], str],
) -> tuple[pd.DataFrame, list[dict]]:
    mapped: dict[str, pd.Series] = {}
    unmapped_actionable: list[dict] = []
    for original_header in frame.columns:
        normalized_header = normalize_metric_header(original_header)
        canonical = dictionary_map.get((family, normalized_header))
        series = frame[original_header]
        if canonical is None:
            if series_has_nonempty_values(series) and normalized_header != "unnamed":
                unmapped_actionable.append(
                    {
                        "component_family": family,
                        "original_header": original_header,
                        "normalized_header": normalized_header,
                    }
                )
            continue
        if canonical in mapped:
            mapped[canonical] = combine_duplicate_series(mapped[canonical], series, label=f"{family}:{canonical}")
        else:
            mapped[canonical] = series.copy()
    return pd.DataFrame(mapped), unmapped_actionable


def component_prefix(family: str) -> str:
    return COMPONENT_FAMILIES[family]["prefix"]


def prepare_component_panel_frame(
    family: str,
    entry: dict,
    parse_result: WorkbookParseResult,
    dictionary_map: dict[tuple[str, str], str],
) -> tuple[pd.DataFrame, list[dict]]:
    mapped, unmapped_actionable = apply_component_mapping(family, parse_result.frame, dictionary_map)
    if mapped.empty:
        return pd.DataFrame(), unmapped_actionable
    opeid_series = mapped.get("opeid8", pd.Series(dtype="string")).apply(standardize_opeid8).astype("string")
    valid_mask = opeid_series.notna()
    out = pd.DataFrame()
    out["opeid8"] = opeid_series[valid_mask].reset_index(drop=True)
    if out.empty:
        return out, unmapped_actionable
    out["opeid6"] = out["opeid8"].apply(derive_opeid6).astype("string")
    out["award_year"] = pd.Series([entry["award_year"]] * len(out), dtype="string")
    out["award_year_start"] = pd.Series([int(entry["award_year_start"])] * len(out), dtype="Int64")
    out["award_year_end"] = pd.Series([int(entry["award_year_end"])] * len(out), dtype="Int64")

    prefix = component_prefix(family)
    for descriptor in ("school", "state", "zip_code", "school_type"):
        series = mapped.get(descriptor)
        if series is None:
            continue
        series = series[valid_mask].reset_index(drop=True)
        if descriptor == "state":
            out[f"{prefix}{descriptor}"] = format_string_series(series, formatter="upper")
        elif descriptor == "school_type":
            out[f"{prefix}{descriptor}"] = format_string_series(series, formatter="title")
        elif descriptor == "zip_code":
            out[f"{prefix}{descriptor}"] = format_zip_series(series)
        else:
            out[f"{prefix}{descriptor}"] = format_string_series(series)

    measure_columns = [col for col in mapped.columns if col not in {"opeid8", "school", "state", "zip_code", "school_type"}]
    for measure in measure_columns:
        series = mapped[measure][valid_mask].reset_index(drop=True)
        out[f"{prefix}{measure}"] = format_integer_series(series)
    out = out.sort_values(["opeid8", "award_year"]).reset_index(drop=True)
    return out, unmapped_actionable


def panel_file_name(stem: str, entries: Sequence[dict]) -> str:
    starts = sorted(int(entry["award_year_start"]) for entry in entries)
    ends = sorted(int(entry["award_year_end"]) for entry in entries)
    return f"{stem}_{starts[0]}_{ends[-1]}.parquet"


def assert_unique_panel_keys(frame: pd.DataFrame, label: str) -> None:
    duplicate_mask = frame.duplicated(["opeid8", "award_year"], keep=False)
    if duplicate_mask.any():
        raise SystemExit(f"{label} contains duplicate (opeid8, award_year) keys.")


def build_component_panel(root: str | Path | None, family: str) -> Path:
    layout = ensure_data_layout(root)
    dictionary_map = load_dictionary_map(layout)
    entries = load_selected_panel_entries(layout, family=family)
    if not entries:
        raise SystemExit(f"No selected panel entries found for {family}.")
    year_frames: list[pd.DataFrame] = []
    summary_rows: list[dict] = []
    branch_rows: list[dict] = []
    unmapped_rows: list[dict] = []
    for entry in entries:
        workbook_path = Path(entry["local_path"])
        parse_result = parse_selected_sheet(workbook_path, family, entry["period_type"])
        panel_frame, unmapped = prepare_component_panel_frame(family, entry, parse_result, dictionary_map)
        unmapped_rows.extend(
            [
                {
                    **row,
                    "filename": entry["filename"],
                    "award_year": entry["award_year"],
                }
                for row in unmapped
            ]
        )
        if panel_frame.empty:
            continue
        assert_unique_panel_keys(panel_frame, f"{family} {entry['award_year']}")
        year_path = component_cross_sections_dir(layout, family) / f"panel_{family}_{int(entry['award_year_start'])}.parquet"
        panel_frame.to_parquet(year_path, index=False)
        year_frames.append(panel_frame)
        branch_count = int((panel_frame["opeid8"].astype("string").str[-2:] != "00").sum())
        branch_rows.append(
            {
                "family": family,
                "award_year": entry["award_year"],
                "rows": int(len(panel_frame)),
                "branch_opeid_rows": branch_count,
            }
        )
        summary_rows.append(
            {
                "family": family,
                "award_year": entry["award_year"],
                "rows": int(len(panel_frame)),
                "columns": int(len(panel_frame.columns)),
                "branch_opeid_rows": branch_count,
            }
        )
    if unmapped_rows:
        pd.DataFrame(unmapped_rows).to_csv(
            layout.checks / "panel_qc" / f"{family}_unmapped_actionable_headers.csv",
            index=False,
        )
        raise SystemExit(f"Unmapped actionable headers encountered while building {family}.")
    stitched = pd.concat(year_frames, ignore_index=True) if year_frames else pd.DataFrame()
    if stitched.empty:
        raise SystemExit(f"No stitched rows produced for {family}.")
    assert_unique_panel_keys(stitched, f"{family} stitched panel")
    if family == "grants":
        filename = panel_file_name("panel_grant_volume", entries)
    elif family == "campus_based":
        filename = panel_file_name("panel_campus_based_volume", entries)
    elif family == "direct_loans":
        filename = panel_file_name("panel_direct_loan_volume", entries)
    else:
        filename = panel_file_name("panel_ffel_loan_volume", entries)
    output_path = component_panel_dir(layout, family) / filename
    stitched.to_parquet(output_path, index=False)
    write_rows(
        layout.checks / "panel_qc" / f"{family}_year_summary.csv",
        summary_rows,
        ["family", "award_year", "rows", "columns", "branch_opeid_rows"],
    )
    write_rows(
        layout.checks / "panel_qc" / f"{family}_branch_opeid_summary.csv",
        branch_rows,
        ["family", "award_year", "rows", "branch_opeid_rows"],
    )
    return output_path


def locate_component_panel(layout: DataRootLayout, family: str) -> Path:
    patterns = {
        "grants": "panel_grant_volume_*.parquet",
        "campus_based": "panel_campus_based_volume_*.parquet",
        "direct_loans": "panel_direct_loan_volume_*.parquet",
        "ffel": "panel_ffel_loan_volume_*.parquet",
        "loans": "panel_loan_volume_*.parquet",
    }
    matches = sorted(component_panel_dir(layout, family).glob(patterns[family]))
    if not matches:
        raise SystemExit(f"Unable to locate component panel for {family}.")
    return matches[-1]


def coalesce_duplicate_metadata_columns(frame: pd.DataFrame, column_name: str) -> pd.DataFrame:
    left = f"{column_name}_x"
    right = f"{column_name}_y"
    if left in frame.columns and right in frame.columns:
        frame[column_name] = frame[left].combine_first(frame[right])
        frame = frame.drop(columns=[left, right])
    elif left in frame.columns:
        frame = frame.rename(columns={left: column_name})
    elif right in frame.columns:
        frame = frame.rename(columns={right: column_name})
    return frame


def outer_merge_panels(left: pd.DataFrame, right: pd.DataFrame) -> pd.DataFrame:
    merged = left.merge(right, on=["opeid8", "award_year"], how="outer")
    for column in ("opeid6", "award_year_start", "award_year_end"):
        merged = coalesce_duplicate_metadata_columns(merged, column)
    return merged


def merge_loan_panels(root: str | Path | None = None) -> Path:
    layout = ensure_data_layout(root)
    direct = pd.read_parquet(locate_component_panel(layout, "direct_loans"))
    ffel = pd.read_parquet(locate_component_panel(layout, "ffel"))
    merged = outer_merge_panels(direct, ffel)
    assert_unique_panel_keys(merged, "merged loan panel")
    entries = load_selected_panel_entries(layout, family="direct_loans") + load_selected_panel_entries(layout, family="ffel")
    output_path = component_panel_dir(layout, "direct_loans") / panel_file_name("panel_loan_volume", entries)
    merged.to_parquet(output_path, index=False)
    return output_path


def coalesce_columns(frame: pd.DataFrame, columns: Sequence[str], *, formatter: str | None = None) -> pd.Series:
    series = pd.Series([pd.NA] * len(frame), dtype="string")
    for column in columns:
        if column not in frame.columns:
            continue
        candidate = frame[column]
        if formatter == "state":
            candidate = format_string_series(candidate, formatter="upper")
        elif formatter == "school_type":
            candidate = format_string_series(candidate, formatter="title")
        elif formatter == "zip_code":
            candidate = format_zip_series(candidate)
        else:
            candidate = format_string_series(candidate)
        series = series.combine_first(candidate)
    return series


FINAL_DESCRIPTOR_COLUMNS = {
    "school": ["grant__school", "campus__school", "loan_direct__school", "loan_ffel__school"],
    "state": ["grant__state", "campus__state", "loan_direct__state", "loan_ffel__state"],
    "zip_code": ["grant__zip_code", "campus__zip_code", "loan_direct__zip_code", "loan_ffel__zip_code"],
    "school_type": ["grant__school_type", "campus__school_type", "loan_direct__school_type", "loan_ffel__school_type"],
}

FINAL_DESCRIPTOR_SOURCE_ALIASES = ["grant_value", "campus_value", "loan_direct_value", "loan_ffel_value"]


def unique_preserving_order(values: Sequence[str]) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        ordered.append(value)
    return ordered


def basic_clean_text(value: object) -> str | None:
    text = safe_text(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text or None


def format_zip_value(value: object) -> str | None:
    text = basic_clean_text(value)
    if text is None:
        return None
    digits = re.sub(r"\D", "", text)
    if not digits:
        return None
    if len(digits) >= 9:
        digits = digits[:9]
        return f"{digits[:5]}-{digits[5:]}"
    return digits[:5].zfill(5)


def zip_base5(value: object) -> str | None:
    formatted = format_zip_value(value)
    if formatted is None:
        return None
    digits = re.sub(r"\D", "", formatted)
    return digits[:5] if digits else None


def school_compare_key(value: object) -> str | None:
    text = basic_clean_text(value)
    if text is None:
        return None
    normalized = text.upper().replace("&", " AND ")
    normalized = re.sub(r"^THE\s+", "", normalized)
    normalized = re.sub(r"[^A-Z0-9]+", " ", normalized)
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return normalized or None


def school_display_score(value: str) -> tuple[int, int, int, int]:
    text = basic_clean_text(value) or ""
    has_lower = int(any(char.islower() for char in text))
    no_leading_the = int(not text.upper().startswith("THE "))
    no_all_caps = int(not text.isupper())
    return (has_lower, no_all_caps, no_leading_the, len(text))


def choose_best_school_display(values: Sequence[str]) -> str | None:
    cleaned = [basic_clean_text(value) for value in values]
    cleaned = [value for value in cleaned if value is not None]
    if not cleaned:
        return None
    indexed = list(enumerate(cleaned))
    best_index, best_value = max(indexed, key=lambda item: (school_display_score(item[1]), -item[0]))
    return best_value


def canonical_school_type(value: object) -> tuple[str | None, str | None]:
    text = basic_clean_text(value)
    if text is None:
        return None, None
    normalized = re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()
    if not normalized:
        return None, None
    foreign = "FOREIGN" in normalized
    if "OTHER" in normalized:
        return "other", "Other"
    if "PUBLIC" in normalized:
        return ("foreign_public", "Foreign Public") if foreign else ("public", "Public")
    if "PROPRIETARY" in normalized or "FOR PROFIT" in normalized:
        return ("foreign_for_profit", "Foreign For-Profit") if foreign else ("private_for_profit", "Private/For-Profit")
    if "PRIVATE" in normalized or "NONPROFIT" in normalized or "NON PROFIT" in normalized:
        return ("foreign_private", "Foreign Private") if foreign else ("private_nonprofit", "Private/Non-Profit")
    if foreign:
        return "foreign", "Foreign"
    return normalize_token(text) or None, text


def resolve_descriptor_values(label: str, ordered_values: Sequence[object]) -> dict:
    if label == "school":
        raw_values = unique_preserving_order([value for value in (basic_clean_text(v) for v in ordered_values) if value])
        normalized_values = unique_preserving_order([value for value in (school_compare_key(v) for v in raw_values) if value])
        if not raw_values:
            return {"clean_value": pd.NA, "raw_values": [], "normalized_values": [], "resolution_status": "missing", "needs_manual_review": False}
        nonempty_count = len([v for v in ordered_values if basic_clean_text(v)])
        if len(normalized_values) == 1:
            clean_value = choose_best_school_display(raw_values)
            if nonempty_count == 1:
                status = "single_source"
            elif len(raw_values) == 1:
                status = "consistent"
            else:
                status = "cosmetic_normalized"
            return {
                "clean_value": clean_value,
                "raw_values": raw_values,
                "normalized_values": normalized_values,
                "resolution_status": status,
                "needs_manual_review": False,
            }
        return {
            "clean_value": raw_values[0],
            "raw_values": raw_values,
            "normalized_values": normalized_values,
            "resolution_status": "manual_review",
            "needs_manual_review": True,
        }

    if label == "state":
        raw_values = unique_preserving_order([value for value in (basic_clean_text(v) for v in ordered_values) if value])
        raw_values = [value.upper() for value in raw_values]
        if not raw_values:
            return {"clean_value": pd.NA, "raw_values": [], "normalized_values": [], "resolution_status": "missing", "needs_manual_review": False}
        nonempty_count = len([v for v in ordered_values if basic_clean_text(v)])
        if len(raw_values) == 1:
            status = "single_source" if nonempty_count == 1 else "consistent"
            return {
                "clean_value": raw_values[0],
                "raw_values": raw_values,
                "normalized_values": raw_values,
                "resolution_status": status,
                "needs_manual_review": False,
            }
        return {
            "clean_value": raw_values[0],
            "raw_values": raw_values,
            "normalized_values": raw_values,
            "resolution_status": "manual_review",
            "needs_manual_review": True,
        }

    if label == "zip_code":
        raw_values = unique_preserving_order([value for value in (format_zip_value(v) for v in ordered_values) if value])
        normalized_values = unique_preserving_order([value for value in (zip_base5(v) for v in raw_values) if value])
        if not raw_values:
            return {"clean_value": pd.NA, "raw_values": [], "normalized_values": [], "resolution_status": "missing", "needs_manual_review": False}
        nonempty_count = len([v for v in ordered_values if format_zip_value(v)])
        if len(raw_values) == 1:
            status = "single_source" if nonempty_count == 1 else "consistent"
            return {
                "clean_value": raw_values[0],
                "raw_values": raw_values,
                "normalized_values": normalized_values,
                "resolution_status": status,
                "needs_manual_review": False,
            }
        if len(normalized_values) == 1:
            return {
                "clean_value": normalized_values[0],
                "raw_values": raw_values,
                "normalized_values": normalized_values,
                "resolution_status": "zip_base5_collapsed",
                "needs_manual_review": False,
            }
        return {
            "clean_value": raw_values[0],
            "raw_values": raw_values,
            "normalized_values": normalized_values,
            "resolution_status": "manual_review",
            "needs_manual_review": True,
        }

    if label == "school_type":
        raw_values = unique_preserving_order([value for value in (basic_clean_text(v) for v in ordered_values) if value])
        canonical_pairs = [canonical_school_type(value) for value in raw_values]
        normalized_values = unique_preserving_order([display for _, display in canonical_pairs if display])
        normalized_keys = unique_preserving_order([key for key, _ in canonical_pairs if key])
        if not raw_values:
            return {"clean_value": pd.NA, "raw_values": [], "normalized_values": [], "resolution_status": "missing", "needs_manual_review": False}
        nonempty_count = len([v for v in ordered_values if basic_clean_text(v)])
        if len(normalized_keys) == 1:
            clean_value = normalized_values[0] if normalized_values else raw_values[0]
            if nonempty_count == 1 and clean_value == raw_values[0]:
                status = "single_source"
            elif nonempty_count == 1:
                status = "single_source_canonicalized"
            elif len(raw_values) == 1 and clean_value == raw_values[0]:
                status = "consistent"
            else:
                status = "canonicalized"
            return {
                "clean_value": clean_value,
                "raw_values": raw_values,
                "normalized_values": normalized_values,
                "resolution_status": status,
                "needs_manual_review": False,
            }
        return {
            "clean_value": normalized_values[0] if normalized_values else raw_values[0],
            "raw_values": raw_values,
            "normalized_values": normalized_values,
            "resolution_status": "manual_review",
            "needs_manual_review": True,
        }

    raise ValueError(f"Unsupported descriptor label: {label}")


def audit_descriptor_columns(frame: pd.DataFrame, descriptor: str, columns: Sequence[str]) -> tuple[pd.Series, pd.DataFrame, pd.DataFrame, list[dict]]:
    present_columns = [column for column in columns if column in frame.columns]
    clean_values: list[object] = []
    detail_rows: list[dict] = []
    manual_rows: list[dict] = []
    summary_counts: dict[tuple[str, bool], int] = {}

    for row in frame[["opeid8", "award_year", *present_columns]].itertuples(index=False, name=None):
        opeid8 = row[0]
        award_year = row[1]
        source_values = list(row[2:])
        result = resolve_descriptor_values(descriptor, source_values)
        clean_values.append(result["clean_value"])
        summary_key = (result["resolution_status"], bool(result["needs_manual_review"]))
        summary_counts[summary_key] = summary_counts.get(summary_key, 0) + 1

        nonempty_count = len(result["raw_values"])
        if nonempty_count <= 1:
            continue
        detail_row = {
            "opeid8": opeid8,
            "award_year": award_year,
            "descriptor": descriptor,
            "resolution_status": result["resolution_status"],
            "needs_manual_review": bool(result["needs_manual_review"]),
            "proposed_clean_value": "" if pd.isna(result["clean_value"]) else str(result["clean_value"]),
            "raw_values": " | ".join(result["raw_values"]),
            "normalized_values": " | ".join(result["normalized_values"]),
        }
        for alias, value in zip(FINAL_DESCRIPTOR_SOURCE_ALIASES, source_values):
            detail_row[alias] = basic_clean_text(value) or ""
        detail_rows.append(detail_row)
        if result["needs_manual_review"]:
            manual_rows.append(detail_row.copy())

    summary_rows = [
        {
            "descriptor": descriptor,
            "resolution_status": status,
            "needs_manual_review": needs_manual_review,
            "rows": count,
        }
        for (status, needs_manual_review), count in sorted(summary_counts.items(), key=lambda item: (item[0][0], item[0][1]))
    ]
    clean_series = pd.Series(clean_values, dtype="string")
    detail = pd.DataFrame(detail_rows)
    manual = pd.DataFrame(manual_rows)
    return clean_series, detail, manual, summary_rows


def descriptor_conflicts(frame: pd.DataFrame, raw_descriptor_columns: Sequence[str], label: str) -> pd.DataFrame:
    rows: list[dict] = []
    for _, row in frame.iterrows():
        values = []
        for column in raw_descriptor_columns:
            if column not in frame.columns:
                continue
            value = row[column]
            if pd.notna(value) and str(value).strip():
                values.append(str(value).strip())
        unique_values = sorted(set(values))
        if len(unique_values) > 1:
            rows.append(
                {
                    "opeid8": row["opeid8"],
                    "award_year": row["award_year"],
                    "descriptor": label,
                    "values": " | ".join(unique_values),
                }
            )
    return pd.DataFrame(rows)


def merge_final_panels(root: str | Path | None = None) -> tuple[Path, Path]:
    layout = ensure_data_layout(root)
    grant = pd.read_parquet(locate_component_panel(layout, "grants"))
    campus = pd.read_parquet(locate_component_panel(layout, "campus_based"))
    loans = pd.read_parquet(locate_component_panel(layout, "loans"))
    raw = outer_merge_panels(outer_merge_panels(grant, campus), loans)
    assert_unique_panel_keys(raw, "final raw panel")
    span_entries = (
        load_selected_panel_entries(layout, family="grants")
        + load_selected_panel_entries(layout, family="campus_based")
        + load_selected_panel_entries(layout, family="direct_loans")
        + load_selected_panel_entries(layout, family="ffel")
    )
    raw_path = layout.panels / "final" / panel_file_name("fsa_volume_reports_raw", span_entries)
    raw.to_parquet(raw_path, index=False)

    clean = raw.copy()
    resolution_detail_frames: list[pd.DataFrame] = []
    manual_review_frames: list[pd.DataFrame] = []
    resolution_summary_rows: list[dict] = []
    for descriptor, columns in FINAL_DESCRIPTOR_COLUMNS.items():
        resolved, detail, manual, summary_rows = audit_descriptor_columns(clean, descriptor, columns)
        clean[descriptor] = resolved
        if not detail.empty:
            resolution_detail_frames.append(detail)
        if not manual.empty:
            manual_review_frames.append(manual)
        resolution_summary_rows.extend(summary_rows)
    assert_unique_panel_keys(clean, "final clean panel")
    clean_path = layout.panels / "final" / panel_file_name("fsa_volume_reports_clean", span_entries)
    clean.to_parquet(clean_path, index=False)

    conflict_frames = [
        descriptor_conflicts(raw, FINAL_DESCRIPTOR_COLUMNS["school"], "school"),
        descriptor_conflicts(raw, FINAL_DESCRIPTOR_COLUMNS["state"], "state"),
        descriptor_conflicts(raw, FINAL_DESCRIPTOR_COLUMNS["zip_code"], "zip_code"),
        descriptor_conflicts(raw, FINAL_DESCRIPTOR_COLUMNS["school_type"], "school_type"),
    ]
    conflicts = pd.concat([frame for frame in conflict_frames if not frame.empty], ignore_index=True) if any(
        not frame.empty for frame in conflict_frames
    ) else pd.DataFrame(columns=["opeid8", "award_year", "descriptor", "values"])
    conflicts.to_csv(layout.checks / "panel_qc" / "final_descriptor_conflicts.csv", index=False)
    resolution_detail = (
        pd.concat(resolution_detail_frames, ignore_index=True)
        if resolution_detail_frames
        else pd.DataFrame(
            columns=[
                "opeid8",
                "award_year",
                "descriptor",
                "resolution_status",
                "needs_manual_review",
                "proposed_clean_value",
                "raw_values",
                "normalized_values",
                *FINAL_DESCRIPTOR_SOURCE_ALIASES,
            ]
        )
    )
    resolution_detail.to_csv(layout.checks / "panel_qc" / "final_descriptor_resolution_detail.csv", index=False)
    manual_review = (
        pd.concat(manual_review_frames, ignore_index=True)
        if manual_review_frames
        else pd.DataFrame(columns=resolution_detail.columns)
    )
    manual_review.to_csv(layout.checks / "panel_qc" / "final_descriptor_manual_review.csv", index=False)
    pd.DataFrame(resolution_summary_rows).sort_values(["descriptor", "resolution_status"]).to_csv(
        layout.checks / "panel_qc" / "final_descriptor_resolution_summary.csv",
        index=False,
    )
    return raw_path, clean_path


def build_panel_dictionary(
    root: str | Path | None = None,
    *,
    input_path: str | Path | None = None,
    output_csv: str | Path | None = None,
    output_parquet: str | Path | None = None,
) -> tuple[Path, Path]:
    layout = ensure_data_layout(root)
    if input_path:
        panel_path = Path(input_path)
    else:
        matches = sorted((layout.panels / "final").glob("fsa_volume_reports_clean_*.parquet"))
        if not matches:
            raise SystemExit("No final clean panel found for panel dictionary build.")
        panel_path = matches[-1]
    dictionary = pd.read_parquet(layout.dictionary / "fsa_volume_dictionary.parquet")
    panel = pd.read_parquet(panel_path)
    rows = []
    for column in panel.columns:
        if column in {"opeid8", "opeid6", "award_year", "award_year_start", "award_year_end", "school", "state", "zip_code", "school_type"}:
            rows.append(
                {
                    "panel_column": column,
                    "component_family": "synthetic" if column in {"school", "state", "zip_code", "school_type"} else "key",
                    "canonical_column": column,
                    "mapping_status": "synthetic_coalesced" if column in {"school", "state", "zip_code", "school_type"} else "panel_key",
                    "source_filename": "",
                    "sheet_name": "",
                    "first_year": int(panel["award_year_start"].min()) if "award_year_start" in panel.columns else pd.NA,
                    "last_year": int(panel["award_year_end"].max()) if "award_year_end" in panel.columns else pd.NA,
                }
            )
            continue
    # Build rows deterministically without relying on fragile vector matching.
    mapped_rows = []
    for _, row in dictionary.iterrows():
        canonical = safe_text(row["canonical_column"])
        if not canonical:
            continue
        for family in COMPONENT_FAMILIES:
            prefix = component_prefix(family)
            panel_column = f"{prefix}{canonical}"
            if panel_column in panel.columns:
                mapped_rows.append(
                    {
                        "panel_column": panel_column,
                        **row.to_dict(),
                    }
                )
    rows.extend(mapped_rows)
    out = pd.DataFrame(rows).drop_duplicates().sort_values(["panel_column", "component_family"]).reset_index(drop=True)
    csv_path = Path(output_csv) if output_csv else layout.dictionary / "fsa_volume_panel_dictionary.csv"
    parquet_path = Path(output_parquet) if output_parquet else layout.dictionary / "fsa_volume_panel_dictionary.parquet"
    out.to_csv(csv_path, index=False)
    out.to_parquet(parquet_path, index=False)
    return csv_path, parquet_path


REVIEW_DESCRIPTOR_ORDER = {"school": 0, "zip_code": 1, "school_type": 2, "state": 3}
REVIEW_SHEET_NAMES = {
    "summary": "Summary",
    "all": "All_Manual_Review",
    "school": "School",
    "zip_code": "Zip_Code",
    "school_type": "School_Type",
    "state": "State",
    "priority_all": "Priority_All",
    "priority_state": "State_Conflicts",
    "priority_renames": "Likely_Renames",
}
PRIORITY_REVIEW_BUCKETS = ("state_mismatch", "likely_rename_pattern")
PRIORITY_REVIEW_BUCKET_ORDER = {
    "state_mismatch": 0,
    "likely_rename_pattern": 1,
}
REVIEW_DECISION_DEFAULTS = {
    "likely_rename_pattern": "confirm_and_accept_proposed_clean_value",
    "substantive_name_conflict": "manual_name_review_required",
    "likely_address_change_or_typo": "verify_best_zip_before_accepting",
    "substantive_zip_conflict": "manual_zip_review_required",
    "likely_label_scheme_conflict": "confirm_and_accept_canonical_school_type",
    "substantive_sector_conflict": "manual_school_type_review_required",
    "state_mismatch": "verify_true_state_conflict",
    "manual_review": "manual_review_required",
}


def manual_review_package_dir(layout: DataRootLayout) -> Path:
    return layout.checks / "panel_qc" / "manual_review_package"


def school_similarity_metrics(values: Sequence[str]) -> tuple[float, bool]:
    normalized = [school_compare_key(value) for value in values]
    normalized = [value for value in normalized if value]
    if len(normalized) < 2:
        return 0.0, False
    token_sets = []
    stopwords = {"THE", "OF", "AT", "IN", "AND", "FOR"}
    for value in normalized:
        tokens = {token for token in value.split() if token and token not in stopwords}
        token_sets.append(tokens)
    max_jaccard = 0.0
    contains = False
    for idx, left in enumerate(normalized):
        for jdx in range(idx + 1, len(normalized)):
            right = normalized[jdx]
            if left in right or right in left:
                contains = True
            union = token_sets[idx] | token_sets[jdx]
            if union:
                score = len(token_sets[idx] & token_sets[jdx]) / len(union)
                max_jaccard = max(max_jaccard, score)
    return max_jaccard, contains


def zip_similarity_metrics(values: Sequence[str]) -> tuple[bool, int | None]:
    normalized = [zip_base5(value) for value in values]
    normalized = [value for value in normalized if value]
    if len(normalized) < 2:
        return False, None
    same_prefix3 = any(left[:3] == right[:3] for idx, left in enumerate(normalized) for right in normalized[idx + 1 :])
    digit_distance: int | None = None
    for idx, left in enumerate(normalized):
        for right in normalized[idx + 1 :]:
            if len(left) != len(right):
                continue
            distance = sum(a != b for a, b in zip(left, right))
            digit_distance = distance if digit_distance is None else min(digit_distance, distance)
    return same_prefix3, digit_distance


def school_type_broad_groups(values: Sequence[str]) -> list[str]:
    groups: list[str] = []
    for value in values:
        key, _ = canonical_school_type(value)
        if key is None:
            continue
        if key.startswith("foreign_"):
            groups.append("foreign")
        elif key.startswith("private_"):
            groups.append("private")
        else:
            groups.append(key)
    return groups


def annotate_manual_review_rows(frame: pd.DataFrame) -> pd.DataFrame:
    out = frame.copy()
    out["descriptor_order"] = out["descriptor"].map(REVIEW_DESCRIPTOR_ORDER).fillna(99).astype(int)
    out["raw_value_count"] = out["raw_values"].fillna("").apply(lambda value: len([item.strip() for item in str(value).split("|") if item.strip()]))
    out["review_sort_rank"] = 9
    out["review_bucket"] = "manual_review"
    out["review_note"] = ""
    out["review_score"] = 0.0

    for index, row in out.iterrows():
        raw_values = [item.strip() for item in str(row["raw_values"]).split("|") if item.strip()]
        descriptor = row["descriptor"]
        if descriptor == "school":
            similarity, contains = school_similarity_metrics(raw_values)
            likely = contains or similarity >= 0.6
            out.at[index, "review_sort_rank"] = 0 if likely else 1
            out.at[index, "review_bucket"] = "likely_rename_pattern" if likely else "substantive_name_conflict"
            note_parts = []
            if contains:
                note_parts.append("normalized_name_contains_other")
            if similarity:
                note_parts.append(f"token_jaccard={similarity:.3f}")
            out.at[index, "review_note"] = "; ".join(note_parts)
            out.at[index, "review_score"] = similarity
        elif descriptor == "zip_code":
            same_prefix3, digit_distance = zip_similarity_metrics(raw_values)
            likely = same_prefix3 or (digit_distance is not None and digit_distance <= 2)
            out.at[index, "review_sort_rank"] = 0 if likely else 1
            out.at[index, "review_bucket"] = "likely_address_change_or_typo" if likely else "substantive_zip_conflict"
            note_parts = []
            if same_prefix3:
                note_parts.append("shared_zip_prefix3")
            if digit_distance is not None:
                note_parts.append(f"min_digit_distance={digit_distance}")
            out.at[index, "review_note"] = "; ".join(note_parts)
            out.at[index, "review_score"] = float(0 if digit_distance is None else max(0, 5 - digit_distance))
        elif descriptor == "school_type":
            groups = unique_preserving_order(school_type_broad_groups(raw_values))
            likely = len(groups) == 1 and bool(groups)
            out.at[index, "review_sort_rank"] = 0 if likely else 1
            out.at[index, "review_bucket"] = "likely_label_scheme_conflict" if likely else "substantive_sector_conflict"
            out.at[index, "review_note"] = "" if not groups else f"broad_groups={'|'.join(groups)}"
            out.at[index, "review_score"] = float(1 if likely else 0)
        elif descriptor == "state":
            out.at[index, "review_sort_rank"] = 1
            out.at[index, "review_bucket"] = "state_mismatch"
            out.at[index, "review_note"] = "different_state_codes_reported"
            out.at[index, "review_score"] = 0.0

    out["review_decision"] = out["review_bucket"].map(REVIEW_DECISION_DEFAULTS).fillna("manual_review_required")
    out["priority_sort_rank"] = out["review_bucket"].map(PRIORITY_REVIEW_BUCKET_ORDER).fillna(9).astype(int)
    out["priority_scope"] = out["review_bucket"].apply(
        lambda value: "highest_priority_review" if value in PRIORITY_REVIEW_BUCKETS else "standard_manual_review"
    )
    sort_columns = ["descriptor_order", "review_sort_rank", "review_score", "opeid8", "award_year"]
    ascending = [True, True, False, True, True]
    out = out.sort_values(sort_columns, ascending=ascending).reset_index(drop=True)
    return out


def build_priority_manual_review_frame(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        return frame.copy()
    priority = frame.loc[frame["review_bucket"].isin(PRIORITY_REVIEW_BUCKETS)].copy()
    priority["priority_group"] = priority["review_bucket"].map(
        {
            "state_mismatch": "true_state_conflict",
            "likely_rename_pattern": "likely_rename",
        }
    ).fillna("other")
    sort_columns = ["priority_sort_rank", "review_score", "opeid8", "award_year"]
    ascending = [True, False, True, True]
    priority = priority.sort_values(sort_columns, ascending=ascending).reset_index(drop=True)
    return priority


def style_review_sheet(worksheet, frame: pd.DataFrame) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 24
    for idx, column_name in enumerate(frame.columns, start=1):
        values = [column_name] + frame[column_name].astype(str).tolist()
        width = min(max(len(str(value)) for value in values) + 2, 48)
        worksheet.column_dimensions[get_column_letter(idx)].width = width


def write_frame_to_sheet(worksheet, frame: pd.DataFrame) -> None:
    worksheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        worksheet.append(list(row))
    style_review_sheet(worksheet, frame)


def build_manual_review_workbook(
    root: str | Path | None = None,
    *,
    input_csv: str | Path | None = None,
    summary_csv: str | Path | None = None,
    output_xlsx: str | Path | None = None,
    output_dir: str | Path | None = None,
) -> tuple[Path, Path]:
    layout = ensure_data_layout(root)
    manual_review_path = Path(input_csv) if input_csv else layout.checks / "panel_qc" / "final_descriptor_manual_review.csv"
    summary_path = Path(summary_csv) if summary_csv else layout.checks / "panel_qc" / "final_descriptor_resolution_summary.csv"
    if not manual_review_path.exists():
        raise SystemExit(f"Missing manual review artifact: {manual_review_path}")
    if not summary_path.exists():
        raise SystemExit(f"Missing descriptor resolution summary: {summary_path}")

    manual = pd.read_csv(manual_review_path, dtype=str).fillna("")
    summary = pd.read_csv(summary_path, dtype={"descriptor": str, "resolution_status": str, "needs_manual_review": bool, "rows": int})
    annotated = annotate_manual_review_rows(manual) if len(manual.columns) else manual.copy()
    priority = build_priority_manual_review_frame(annotated) if "review_bucket" in annotated.columns else annotated.copy()

    target_package_dir = Path(output_dir) if output_dir else manual_review_package_dir(layout)
    use_atomic_replace = output_xlsx is None and target_package_dir.exists()
    if use_atomic_replace:
        package_dir = target_package_dir.parent / f".{target_package_dir.name}__staging"
        if package_dir.exists():
            shutil.rmtree(package_dir)
    else:
        package_dir = target_package_dir
    package_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = Path(output_xlsx) if output_xlsx else package_dir / "final_descriptor_manual_review_workbook.xlsx"
    priority_workbook_path = package_dir / "priority_manual_review_workbook.xlsx"

    annotated.to_csv(package_dir / "all_manual_review.csv", index=False)
    for descriptor in REVIEW_DESCRIPTOR_ORDER:
        descriptor_frame = annotated.loc[annotated["descriptor"] == descriptor].copy()
        descriptor_frame.to_csv(package_dir / f"{descriptor}_manual_review.csv", index=False)
    priority.to_csv(package_dir / "priority_manual_review.csv", index=False)
    priority.loc[priority["review_bucket"] == "state_mismatch"].copy().to_csv(package_dir / "priority_state_conflicts.csv", index=False)
    priority.loc[priority["review_bucket"] == "likely_rename_pattern"].copy().to_csv(package_dir / "priority_likely_renames.csv", index=False)

    workbook = Workbook()
    summary_ws = workbook.active
    summary_ws.title = REVIEW_SHEET_NAMES["summary"]

    summary_intro = pd.DataFrame(
        [
            {"item": "manual_review_rows", "value": int(len(annotated))},
            {"item": "package_dir", "value": str(package_dir)},
            {"item": "source_manual_review_csv", "value": str(manual_review_path)},
            {"item": "source_resolution_summary_csv", "value": str(summary_path)},
        ]
    )
    write_frame_to_sheet(summary_ws, summary_intro)
    summary_ws.append([])

    summary_table = summary.copy()
    summary_table["needs_manual_review"] = summary_table["needs_manual_review"].astype(bool)
    for row in summary_table.itertuples(index=False, name=None):
        summary_ws.append(list(row))
    summary_start_row = len(summary_intro) + 3
    summary_ws.auto_filter.ref = f"A{summary_start_row}:D{summary_start_row + len(summary_table)}"
    for cell in summary_ws[summary_start_row]:
        cell.fill = PatternFill(fill_type="solid", fgColor="FDE9D9")
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for idx in range(1, 5):
        width = min(max(len(str(summary_ws.cell(row=row_idx, column=idx).value or "")) for row_idx in range(1, summary_ws.max_row + 1)) + 2, 48)
        summary_ws.column_dimensions[get_column_letter(idx)].width = width

    if not annotated.empty:
        all_sheet = workbook.create_sheet(REVIEW_SHEET_NAMES["all"])
        write_frame_to_sheet(all_sheet, annotated)
        for descriptor in REVIEW_DESCRIPTOR_ORDER:
            descriptor_frame = annotated.loc[annotated["descriptor"] == descriptor].copy()
            sheet = workbook.create_sheet(REVIEW_SHEET_NAMES[descriptor])
            write_frame_to_sheet(sheet, descriptor_frame if not descriptor_frame.empty else pd.DataFrame(columns=annotated.columns))

    workbook.save(workbook_path)

    priority_workbook = Workbook()
    priority_summary_ws = priority_workbook.active
    priority_summary_ws.title = REVIEW_SHEET_NAMES["summary"]
    priority_summary = pd.DataFrame(
        [
            {"item": "priority_manual_review_rows", "value": int(len(priority))},
            {"item": "true_state_conflict_rows", "value": int((priority["review_bucket"] == "state_mismatch").sum()) if not priority.empty else 0},
            {"item": "likely_rename_rows", "value": int((priority["review_bucket"] == "likely_rename_pattern").sum()) if not priority.empty else 0},
            {"item": "default_state_review_decision", "value": REVIEW_DECISION_DEFAULTS["state_mismatch"]},
            {"item": "default_likely_rename_review_decision", "value": REVIEW_DECISION_DEFAULTS["likely_rename_pattern"]},
            {"item": "source_manual_review_csv", "value": str(manual_review_path)},
        ]
    )
    write_frame_to_sheet(priority_summary_ws, priority_summary)

    priority_columns = priority.columns if not priority.empty else annotated.columns
    priority_all_sheet = priority_workbook.create_sheet(REVIEW_SHEET_NAMES["priority_all"])
    write_frame_to_sheet(
        priority_all_sheet,
        priority if not priority.empty else pd.DataFrame(columns=priority_columns),
    )
    priority_state = priority.loc[priority["review_bucket"] == "state_mismatch"].copy() if not priority.empty else pd.DataFrame(columns=priority_columns)
    priority_state_sheet = priority_workbook.create_sheet(REVIEW_SHEET_NAMES["priority_state"])
    write_frame_to_sheet(priority_state_sheet, priority_state if not priority_state.empty else pd.DataFrame(columns=priority_columns))
    priority_renames = priority.loc[priority["review_bucket"] == "likely_rename_pattern"].copy() if not priority.empty else pd.DataFrame(columns=priority_columns)
    priority_rename_sheet = priority_workbook.create_sheet(REVIEW_SHEET_NAMES["priority_renames"])
    write_frame_to_sheet(priority_rename_sheet, priority_renames if not priority_renames.empty else pd.DataFrame(columns=priority_columns))

    priority_workbook.save(priority_workbook_path)
    if use_atomic_replace:
        backup_suffix = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        backup_dir = target_package_dir.parent / f"{target_package_dir.name}__backup_{backup_suffix}"
        shutil.move(str(target_package_dir), str(backup_dir))
        shutil.move(str(package_dir), str(target_package_dir))
        package_dir = target_package_dir
        workbook_path = target_package_dir / "final_descriptor_manual_review_workbook.xlsx"
    return workbook_path, package_dir


def source_qaqc(root: str | Path | None = None) -> pd.DataFrame:
    layout = ensure_data_layout(root)
    selected_path = selected_panel_files_path(layout)
    profile_path = layout.checks / "source_qc" / "workbook_profiles.csv"
    results: list[dict] = []
    selected_exists = selected_path.exists()
    results.append({"check": "selected_panel_files_exists", "passed": selected_exists, "details": str(selected_path)})
    if selected_exists:
        selected = pd.read_csv(selected_path, dtype=str).fillna("")
        missing_files = selected[~selected["local_path"].apply(lambda p: Path(p).exists())]
        results.append(
            {
                "check": "selected_files_exist",
                "passed": missing_files.empty,
                "details": "" if missing_files.empty else f"missing={len(missing_files)}",
            }
        )
    else:
        selected = pd.DataFrame()
        results.append({"check": "selected_files_exist", "passed": False, "details": "selected_panel_files.csv missing"})
    profile_exists = profile_path.exists()
    results.append({"check": "workbook_profiles_exists", "passed": profile_exists, "details": str(profile_path)})
    if profile_exists and not selected.empty:
        profiles = pd.read_csv(profile_path, dtype=str).fillna("")
        profiled = selected.merge(
            profiles[["family", "filename", "selected_sheet", "parse_warnings_json"]],
            on=["family", "filename"],
            how="left",
        )
        missing_sheet = profiled["selected_sheet"].eq("")
        results.append(
            {
                "check": "selected_sheet_found_for_selected_files",
                "passed": not bool(missing_sheet.any()),
                "details": "" if not bool(missing_sheet.any()) else f"missing={int(missing_sheet.sum())}",
            }
        )
        year_gap_rows = []
        for family, group in selected.groupby("family"):
            years = sorted(group["award_year_start"].astype(int).unique().tolist())
            expected = list(range(years[0], years[-1] + 1))
            gaps = [year for year in expected if year not in years]
            if gaps:
                year_gap_rows.append({"family": family, "gaps": ",".join(str(year) for year in gaps)})
        results.append(
            {
                "check": "selected_years_contiguous_within_family",
                "passed": len(year_gap_rows) == 0,
                "details": "" if not year_gap_rows else json.dumps(year_gap_rows),
            }
        )
    else:
        results.append({"check": "selected_sheet_found_for_selected_files", "passed": False, "details": "profiles missing"})
        results.append({"check": "selected_years_contiguous_within_family", "passed": False, "details": "profiles or selected files missing"})
    out = pd.DataFrame(results)
    out.to_csv(layout.checks / "source_qc" / "source_qaqc_summary.csv", index=False)
    return out


def summarize_panel(path: Path) -> dict:
    frame = pd.read_parquet(path)
    duplicate_keys = int(frame.duplicated(["opeid8", "award_year"]).sum()) if {"opeid8", "award_year"}.issubset(frame.columns) else pd.NA
    return {
        "path": str(path),
        "rows": int(len(frame)),
        "columns": int(len(frame.columns)),
        "duplicate_keys": duplicate_keys,
        "min_award_year_start": int(frame["award_year_start"].min()) if "award_year_start" in frame.columns else pd.NA,
        "max_award_year_end": int(frame["award_year_end"].max()) if "award_year_end" in frame.columns else pd.NA,
    }


def panel_qaqc(root: str | Path | None = None) -> pd.DataFrame:
    layout = ensure_data_layout(root)
    panel_paths = [
        locate_component_panel(layout, "grants"),
        locate_component_panel(layout, "campus_based"),
        locate_component_panel(layout, "direct_loans"),
        locate_component_panel(layout, "ffel"),
        locate_component_panel(layout, "loans"),
    ]
    panel_paths.extend(sorted((layout.panels / "final").glob("fsa_volume_reports_*.parquet")))
    rows = [summarize_panel(path) for path in panel_paths if path.exists()]
    summary = pd.DataFrame(rows)
    summary.to_csv(layout.checks / "panel_qc" / "panel_qaqc_summary.csv", index=False)

    coverage_rows = []
    for path in panel_paths:
        if not path.exists():
            continue
        frame = pd.read_parquet(path)
        if "award_year" not in frame.columns:
            continue
        counts = frame["award_year"].value_counts().sort_index()
        for award_year, count in counts.items():
            coverage_rows.append({"panel": path.name, "award_year": award_year, "rows": int(count)})
    pd.DataFrame(coverage_rows).to_csv(layout.checks / "panel_qc" / "panel_coverage_matrix.csv", index=False)
    return summary


def acceptance_audit(root: str | Path | None = None) -> pd.DataFrame:
    layout = ensure_data_layout(root)
    results: list[dict] = []

    source_summary_path = layout.checks / "source_qc" / "source_qaqc_summary.csv"
    results.append(
        {
            "check": "source_qaqc_summary_exists",
            "passed": source_summary_path.exists(),
            "details": str(source_summary_path),
        }
    )
    if source_summary_path.exists():
        source_summary = pd.read_csv(source_summary_path)
        for row in source_summary.to_dict("records"):
            results.append(
                {
                    "check": f"source::{row['check']}",
                    "passed": bool(row["passed"]),
                    "details": row.get("details", ""),
                }
            )

    dictionary_summary_path = layout.checks / "dictionary_qc" / "dictionary_qaqc_summary.csv"
    results.append(
        {
            "check": "dictionary_qaqc_summary_exists",
            "passed": dictionary_summary_path.exists(),
            "details": str(dictionary_summary_path),
        }
    )
    if dictionary_summary_path.exists():
        dictionary_summary = pd.read_csv(dictionary_summary_path)
        unmapped = int(
            dictionary_summary.loc[dictionary_summary["metric"] == "unmapped_actionable_rows", "value"].astype(int).sum()
        )
        results.append(
            {
                "check": "dictionary_has_no_unmapped_actionable_headers",
                "passed": unmapped == 0,
                "details": f"unmapped_actionable_rows={unmapped}",
            }
        )

    panel_summary_path = layout.checks / "panel_qc" / "panel_qaqc_summary.csv"
    results.append({"check": "panel_qaqc_summary_exists", "passed": panel_summary_path.exists(), "details": str(panel_summary_path)})
    if panel_summary_path.exists():
        panel_summary = pd.read_csv(panel_summary_path)
        duplicate_failures = int(panel_summary["duplicate_keys"].fillna(0).astype(int).sum())
        results.append(
            {
                "check": "no_duplicate_panel_keys",
                "passed": duplicate_failures == 0,
                "details": f"duplicate_keys={duplicate_failures}",
            }
        )

    conflict_path = layout.checks / "panel_qc" / "final_descriptor_conflicts.csv"
    results.append({"check": "descriptor_conflict_artifact_exists", "passed": conflict_path.exists(), "details": str(conflict_path)})
    if conflict_path.exists():
        conflicts = pd.read_csv(conflict_path)
        results.append(
            {
                "check": "descriptor_conflicts_are_auditable",
                "passed": True,
                "details": f"conflict_rows={len(conflicts)}",
            }
        )
    resolution_summary_path = layout.checks / "panel_qc" / "final_descriptor_resolution_summary.csv"
    results.append(
        {
            "check": "descriptor_resolution_summary_exists",
            "passed": resolution_summary_path.exists(),
            "details": str(resolution_summary_path),
        }
    )
    manual_review_path = layout.checks / "panel_qc" / "final_descriptor_manual_review.csv"
    results.append(
        {
            "check": "descriptor_manual_review_artifact_exists",
            "passed": manual_review_path.exists(),
            "details": str(manual_review_path),
        }
    )
    if manual_review_path.exists():
        manual_review = pd.read_csv(manual_review_path)
        results.append(
            {
                "check": "descriptor_manual_review_rows_are_auditable",
                "passed": True,
                "details": f"manual_review_rows={len(manual_review)}",
            }
        )
    workbook_path = manual_review_package_dir(layout) / "final_descriptor_manual_review_workbook.xlsx"
    results.append(
        {
            "check": "descriptor_manual_review_workbook_exists",
            "passed": workbook_path.exists(),
            "details": str(workbook_path),
        }
    )
    priority_workbook_path = manual_review_package_dir(layout) / "priority_manual_review_workbook.xlsx"
    results.append(
        {
            "check": "priority_manual_review_workbook_exists",
            "passed": priority_workbook_path.exists(),
            "details": str(priority_workbook_path),
        }
    )

    selected_path = selected_panel_files_path(layout)
    if selected_path.exists():
        selected = pd.read_csv(selected_path, dtype=str).fillna("")
        family_to_panel = {
            "grants": locate_component_panel(layout, "grants"),
            "campus_based": locate_component_panel(layout, "campus_based"),
            "direct_loans": locate_component_panel(layout, "direct_loans"),
            "ffel": locate_component_panel(layout, "ffel"),
        }
        for family, panel_path in family_to_panel.items():
            panel = pd.read_parquet(panel_path)
            expected_years = sorted(selected.loc[selected["family"] == family, "award_year"].unique().tolist())
            actual_years = sorted(panel["award_year"].dropna().astype(str).unique().tolist())
            results.append(
                {
                    "check": f"{family}::year_coverage_matches_selected_files",
                    "passed": expected_years == actual_years,
                    "details": f"expected={expected_years} actual={actual_years}",
                }
            )

    acceptance = pd.DataFrame(results)
    out_csv = layout.checks / "acceptance_qc" / "acceptance_summary.csv"
    acceptance.to_csv(out_csv, index=False)
    passed = int(acceptance["passed"].astype(bool).sum())
    failed = int((~acceptance["passed"].astype(bool)).sum())
    md_lines = [
        "# Acceptance Audit",
        "",
        f"- Checks passed: `{passed}`",
        f"- Checks failed: `{failed}`",
        "",
        "| Check | Passed | Details |",
        "| --- | --- | --- |",
    ]
    for row in acceptance.to_dict("records"):
        md_lines.append(f"| {row['check']} | {bool(row['passed'])} | {row.get('details', '')} |")
    (layout.checks / "acceptance_qc" / "acceptance_summary.md").write_text("\n".join(md_lines), encoding="utf-8")
    return acceptance
