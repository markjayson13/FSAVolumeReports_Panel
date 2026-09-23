from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from helpers import load_script_module, write_workbook
from openpyxl import load_workbook


utils = load_script_module("fsa_build_utils_descriptor_resolution", "Scripts/fsa_build_utils.py")


class DescriptorResolutionTests(unittest.TestCase):
    def test_school_cosmetic_difference_resolves_without_manual_review(self) -> None:
        result = utils.resolve_descriptor_values(
            "school",
            [
                "ALABAMA AGRICULTURAL & MECHANICAL UNIVERSITY",
                "Alabama Agricultural & Mechanical University",
            ],
        )
        self.assertEqual(result["clean_value"], "Alabama Agricultural & Mechanical University")
        self.assertEqual(result["resolution_status"], "cosmetic_normalized")
        self.assertFalse(result["needs_manual_review"])

    def test_school_substantive_difference_requires_manual_review(self) -> None:
        result = utils.resolve_descriptor_values(
            "school",
            [
                "AUBURN UNIVERSITY",
                "AUBURN UNIVERSITY-AUBURN",
            ],
        )
        self.assertEqual(result["resolution_status"], "manual_review")
        self.assertTrue(result["needs_manual_review"])

    def test_school_punctuation_diacritics_and_parenthetic_article_are_cosmetic(self) -> None:
        for values in (["UNIVERSITY OF TEXAS M.D. ANDERSON CANCER CENTER", "UNIVERSITY OF TEXAS MD ANDERSON CANCER CENTER"],
                       ["MASTER'S COLLEGE AND SEMINARY", "MASTERS COLLEGE & SEMINARY (THE)"],
                       ["LA' JAMES COLLEGE OF HAIRSTYLING", "LA'JAMES COLLEGE OF HAIRSTYLING"],
                       ["COLLEGE OF IDAHO, THE", "COLLEGE OF IDAHO (THE)"],
                       ["Université de Montréal", "Universite de Montreal"]):
            with self.subTest(values=values):
                self.assertFalse(utils.resolve_descriptor_values("school", values)["needs_manual_review"])

    def test_name_qualifiers_abbreviations_and_renames_are_not_blindly_collapsed(self) -> None:
        for values in (["BISHOP STATE COMMUNITY COLLEGE - MAIN CAMPUS", "Bishop State Community College"],
                       ["ST. AUGUSTINE COLLEGE", "Saint Augustine College"],
                       ["EVEREST COLLEGE", "Bryman College"],
                       ["YESHIVA D'MONSEY - (ADMINISTRATIVE OFFICES ONLY)", "Yeshiva D'Monsey"]):
            with self.subTest(values=values):
                self.assertTrue(utils.resolve_descriptor_values("school", values)["needs_manual_review"])

    def test_only_reversible_encoding_damage_is_repaired(self) -> None:
        good = utils.resolve_descriptor_values("school", ["UNIVERSIDAD ANA G. MÉNDEZ", "Universidad Ana G. MÃ©ndez"])
        self.assertFalse(good["needs_manual_review"])
        self.assertEqual(good["clean_value"], "Universidad Ana G. Méndez")
        damaged = utils.resolve_descriptor_values("school", ["Universidad del Sagrado Corazón", "Universidad del Sagrado Coraz?n"])
        self.assertTrue(damaged["needs_manual_review"])

    def test_official_annual_hd_postal_examples_match_after_restoring_width(self) -> None:
        # Source-ledger/HD anchor evidence is recorded in Documentation/descriptor_resolution.md.
        # Same full OPEID and AY: Husson 00204300, 2006; Middlesex 00261500, 2013.
        for raw, state, hd_zip in [("44012999", "ME", "04401-2999"), ("88183050", "NJ", "08818-3050")]:
            self.assertEqual(utils.format_zip_value(raw, state=state), hd_zip)

    def test_component_preserves_raw_postal_token_while_normalizing_display(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            headers = ["OPE ID", "School", "State", "Zip Code", "School Type"]
            path = write_workbook(Path(temp) / "postal.xlsx", {"Data": [headers,
                [261500, "MIDDLESEX COUNTY COLLEGE", "NJ", 88183050, "PUBLIC"]]})
            parsed = utils.parse_selected_sheet(path, "direct_loans", "annual")
            canonical = ["opeid8", "school", "state", "zip_code", "school_type"]
            mapping = {("direct_loans", utils.normalize_metric_header(h)): c for h, c in zip(headers, canonical)}
            frame, _ = utils.prepare_component_panel_frame("direct_loans", {
                "award_year": "2013-2014", "award_year_start": 2013, "award_year_end": 2014,
                "filename": path.name}, parsed, mapping)
            self.assertEqual(frame.loc[0, "opeid8"], "00261500")
            self.assertEqual(frame.loc[0, "loan_direct__raw_zip_code"], "88183050")
            self.assertEqual(frame.loc[0, "loan_direct__zip_code"], "08818-3050")

    def test_numeric_zip_plus4_restores_width_only_with_domestic_context(self) -> None:
        self.assertEqual(utils.format_zip_value("88183050", state="NJ"), "08818-3050")
        self.assertEqual(utils.format_zip_value("9610000", state="PR"), "00961-0000")
        self.assertEqual(utils.format_zip_value("601", state="PR"), "00601")
        self.assertEqual(utils.format_zip_value("88183050"), "88183050")
        self.assertEqual(utils.format_zip_value("123456", state="CA"), "123456")

    def test_foreign_postcodes_are_preserved_including_foreign_marker_state_collision(self) -> None:
        self.assertEqual(utils.format_zip_value("WC1E 6BT", school_type="Foreign Public"), "WC1E 6BT")
        self.assertEqual(utils.format_zip_value("H3A 0G4", state="CA", school_type="Foreign Private"), "H3A 0G4")
        self.assertEqual(utils.format_zip_value("88183050", state="NJ", school_type="Foreign Private"), "88183050")
        self.assertEqual(utils.zip_base5("WC1E 6BT"), "WC1E 6BT")

    def test_not_for_profit_is_not_accidentally_for_profit(self) -> None:
        self.assertEqual(utils.canonical_school_type("Private, not-for-profit")[0], "private_nonprofit")
        self.assertEqual(utils.canonical_school_type("Private, for-profit")[0], "private_for_profit")

    def test_school_type_private_variants_canonicalize(self) -> None:
        result = utils.resolve_descriptor_values(
            "school_type",
            [
                "Private",
                "Private/Non-Profit",
                "Private-Nonprofit",
            ],
        )
        self.assertEqual(result["clean_value"], "Private/Non-Profit")
        self.assertEqual(result["resolution_status"], "canonicalized")
        self.assertFalse(result["needs_manual_review"])

    def test_zip_conflict_collapses_to_base5_when_equivalent(self) -> None:
        result = utils.resolve_descriptor_values(
            "zip_code",
            [
                "35762",
                "35762-1357",
            ],
        )
        self.assertEqual(result["clean_value"], "35762")
        self.assertEqual(result["resolution_status"], "zip_base5_collapsed")
        self.assertFalse(result["needs_manual_review"])

    def test_state_conflict_requires_manual_review(self) -> None:
        result = utils.resolve_descriptor_values(
            "state",
            [
                "ca",
                "IA",
            ],
        )
        self.assertEqual(result["clean_value"], "CA")
        self.assertEqual(result["resolution_status"], "manual_review")
        self.assertTrue(result["needs_manual_review"])
        self.assertFalse(utils.resolve_descriptor_values("state", ["ca", "CA"])["needs_manual_review"])

    def test_build_manual_review_workbook_outputs_sorted_descriptor_tabs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "Checks" / "panel_qc").mkdir(parents=True, exist_ok=True)
            manual = utils.pd.DataFrame(
                [
                    {
                        "opeid8": "00100002",
                        "award_year": "2001-2002",
                        "descriptor": "state",
                        "resolution_status": "manual_review",
                        "needs_manual_review": True,
                        "proposed_clean_value": "CA",
                        "raw_values": "CA | IA",
                        "normalized_values": "CA | IA",
                        "grant_value": "CA",
                        "campus_value": "",
                        "loan_direct_value": "",
                        "loan_ffel_value": "IA",
                    },
                    {
                        "opeid8": "00100000",
                        "award_year": "2001-2002",
                        "descriptor": "school",
                        "resolution_status": "manual_review",
                        "needs_manual_review": True,
                        "proposed_clean_value": "AUBURN UNIVERSITY-AUBURN",
                        "raw_values": "AUBURN UNIVERSITY-AUBURN | AUBURN UNIVERSITY",
                        "normalized_values": "AUBURN UNIVERSITY AUBURN | AUBURN UNIVERSITY",
                        "grant_value": "",
                        "campus_value": "AUBURN UNIVERSITY-AUBURN",
                        "loan_direct_value": "",
                        "loan_ffel_value": "AUBURN UNIVERSITY",
                    },
                    {
                        "opeid8": "00100001",
                        "award_year": "2001-2002",
                        "descriptor": "school",
                        "resolution_status": "manual_review",
                        "needs_manual_review": True,
                        "proposed_clean_value": "ALPHA COLLEGE",
                        "raw_values": "ALPHA COLLEGE | BETA INSTITUTE",
                        "normalized_values": "ALPHA COLLEGE | BETA INSTITUTE",
                        "grant_value": "",
                        "campus_value": "ALPHA COLLEGE",
                        "loan_direct_value": "",
                        "loan_ffel_value": "BETA INSTITUTE",
                    },
                ]
            )
            manual_path = root / "Checks" / "panel_qc" / "final_descriptor_manual_review.csv"
            manual.to_csv(manual_path, index=False)
            summary = utils.pd.DataFrame(
                [
                    {"descriptor": "school", "resolution_status": "manual_review", "needs_manual_review": True, "rows": 2},
                    {"descriptor": "state", "resolution_status": "manual_review", "needs_manual_review": True, "rows": 1},
                ]
            )
            summary_path = root / "Checks" / "panel_qc" / "final_descriptor_resolution_summary.csv"
            summary.to_csv(summary_path, index=False)

            workbook_path, package_dir = utils.build_manual_review_workbook(root)
            self.assertTrue(workbook_path.exists())
            self.assertTrue((package_dir / "school_manual_review.csv").exists())
            self.assertTrue((package_dir / "priority_manual_review.csv").exists())
            self.assertTrue((package_dir / "priority_manual_review_workbook.xlsx").exists())

            wb = load_workbook(workbook_path)
            self.assertIn("School", wb.sheetnames)
            school_sheet = wb["School"]
            first_data_row = [school_sheet.cell(row=2, column=col).value for col in range(1, school_sheet.max_column + 1)]
            headers = [school_sheet.cell(row=1, column=col).value for col in range(1, school_sheet.max_column + 1)]
            row_map = dict(zip(headers, first_data_row))
            self.assertEqual(row_map["review_bucket"], "likely_rename_pattern")
            self.assertEqual(row_map["review_decision"], "confirm_and_accept_proposed_clean_value")
            self.assertEqual(row_map["opeid8"], "00100000")

            priority_wb = load_workbook(package_dir / "priority_manual_review_workbook.xlsx")
            self.assertIn("Priority_All", priority_wb.sheetnames)
            self.assertIn("State_Conflicts", priority_wb.sheetnames)
            self.assertIn("Likely_Renames", priority_wb.sheetnames)

            priority_sheet = priority_wb["Priority_All"]
            priority_headers = [priority_sheet.cell(row=1, column=col).value for col in range(1, priority_sheet.max_column + 1)]
            priority_first_row = [priority_sheet.cell(row=2, column=col).value for col in range(1, priority_sheet.max_column + 1)]
            priority_row_map = dict(zip(priority_headers, priority_first_row))
            self.assertEqual(priority_row_map["descriptor"], "state")
            self.assertEqual(priority_row_map["review_bucket"], "state_mismatch")
            self.assertEqual(priority_row_map["review_decision"], "verify_true_state_conflict")

            second_workbook_path, second_package_dir = utils.build_manual_review_workbook(root)
            self.assertEqual(second_workbook_path, second_package_dir / "final_descriptor_manual_review_workbook.xlsx")
            self.assertTrue(second_workbook_path.exists())
            backup_dirs = list((root / "Checks" / "panel_qc").glob("manual_review_package__backup_*"))
            self.assertTrue(backup_dirs)


if __name__ == "__main__":
    unittest.main()
