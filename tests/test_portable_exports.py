"""Adversarial round trips for the portable research release.

The fixture deliberately combines identifier, missingness, and text cases that
spreadsheet/Stata/CSV inference commonly changes without reporting an error.
"""
from __future__ import annotations

import json
from pathlib import Path
import tempfile
import unittest

import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook

from helpers import load_script_module


exports = load_script_module("fsa_portable_exports_test", "Scripts/fsa_portable_exports.py")


def export_fixture() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    columns = [
        ("unitid", "unitid", "numeric", ""),
        ("award_year_start", "award_year_start", "numeric", ""),
        ("grant__opeid8", "gr_opeid8", "string", ""),
        ("grant__school", "gr_school", "string", ""),
        ("grant__pell_disbursements__raw_token", "gr_pell_raw", "string", ""),
        ("grant__pell_disbursements", "gr_pell_amt", "numeric", ""),
        ("grant__pell_recipients", "gr_pell_n", "numeric", ""),
        ("grant__pell_disbursements__status", "gr_pell_status", "coded_category", "status"),
        ("grant__ipeds_in_strict_sample", "gr_strict", "coded_category", "boolean"),
        ("grant__pell_disbursements__upper_bound", "gr_pell_upper", "numeric", ""),
    ]
    panel = pd.DataFrame({
        "unitid": pd.Series([100001, 100002, 100003, 100004, 100005], dtype="Int64"),
        "award_year_start": pd.Series([2024] * 5, dtype="Int64"),
        "grant__opeid8": pd.Series(["00100000", "00100001", "00012300", pd.NA, "99999999"], dtype="string"),
        "grant__school": pd.Series(["École\t東京", "=2+2", "NA", "", pd.NA], dtype="string"),
        "grant__pell_disbursements__raw_token": pd.Series(["x" * 2300, "@SUM(A1)", 'line1\nline2,"quoted"', "", pd.NA], dtype="string"),
        "grant__pell_disbursements": pd.Series([0.0, 0.1, 123456789.12345679, pd.NA, -0.125], dtype="Float64"),
        "grant__pell_recipients": pd.Series([0, 1, 1234567890, pd.NA, 5], dtype="Int64"),
        "grant__pell_disbursements__status": pd.Series(["observed", "observed", "suppressed", "source_blank", pd.NA], dtype="string"),
        "grant__ipeds_in_strict_sample": pd.Series([True, False, pd.NA, True, False], dtype="boolean"),
        "grant__pell_disbursements__upper_bound": pd.Series([pd.NA] * 5, dtype="Float64"),
    })
    metadata = pd.DataFrame([
        {
            "canonical_name": canonical,
            "portable_name": portable,
            "variable_label": f"Research label: {portable}",
            "description": f"Full definition for {canonical}; missing values are not zero.",
            "units": "USD" if portable == "gr_pell_amt" else "text" if storage == "string" else "count/code",
            "role": "identifier" if portable in {"unitid", "gr_opeid8"} else "measure",
            "source_family": "grant",
            "original_dtype": str(panel[canonical].dtype),
            "export_storage": storage,
            "category_kind": category,
        }
        for canonical, portable, storage, category in columns
    ])
    labels = pd.DataFrame([
        {"canonical_name": "grant__pell_disbursements__status", "portable_name": "gr_pell_status", "code": code, "value": value, "label": value, "description": value}
        for code, value in [(1, "observed"), (2, "suppressed"), (3, "source_blank")]
    ] + [
        {"canonical_name": "grant__ipeds_in_strict_sample", "portable_name": "gr_strict", "code": code, "value": value, "label": label, "description": label}
        for code, value, label in [(0, "false", "No"), (1, "true", "Yes")]
    ])
    return panel, metadata, labels


class PortableExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.panel, self.metadata, self.labels = export_fixture()
        self.frame = exports.prepare_export_frame(self.panel, self.metadata, self.labels)
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_preparation_preserves_identifiers_zero_and_nullable_flags(self) -> None:
        self.assertEqual(self.frame["gr_opeid8"].iloc[0], "00100000")
        self.assertEqual(self.frame["gr_school"].iloc[2], "NA")
        self.assertEqual(self.frame["gr_pell_amt"].iloc[0], 0)
        self.assertTrue(pd.isna(self.frame["gr_pell_amt"].iloc[3]))
        self.assertEqual(self.frame["gr_strict"].iloc[0], 1)
        self.assertEqual(self.frame["gr_strict"].iloc[1], 0)
        self.assertTrue(pd.isna(self.frame["gr_strict"].iloc[2]))
        self.assertEqual(self.frame["gr_pell_status"].iloc[2], 2)
        self.assertTrue(self.frame["gr_pell_upper"].isna().all())
        self.assertEqual(self.panel["grant__opeid8"].iloc[0], "00100000")
        self.assertEqual(list(self.frame.columns), self.metadata["portable_name"].tolist())

    def test_unknown_category_is_rejected_instead_of_becoming_missing(self) -> None:
        panel = self.panel.copy()
        panel.loc[0, "grant__pell_disbursements__status"] = "unreviewed_new_status"
        with self.assertRaises(ValueError):
            exports.prepare_export_frame(panel, self.metadata, self.labels)

    def test_duplicate_portable_names_cannot_drop_a_column(self) -> None:
        metadata = self.metadata.copy()
        metadata.loc[1, "portable_name"] = "unitid"
        with self.assertRaises(ValueError):
            exports.prepare_export_frame(self.panel, metadata, self.labels)

    def test_csv_roundtrip_preserves_na_literal_empty_null_unicode_and_double_precision(self) -> None:
        path = self.root / "data.csv.gz"
        exports.write_csv(self.frame, path)
        reread = exports.read_csv(path, self.metadata)
        exports.verify_frame(self.frame, reread, self.metadata, "csv")
        self.assertEqual(reread.loc[0, "gr_opeid8"], "00100000")
        self.assertEqual(reread.loc[2, "gr_school"], "NA")
        self.assertEqual(reread.loc[3, "gr_school"], "")
        self.assertTrue(pd.isna(reread.loc[4, "gr_school"]))
        self.assertEqual(reread.loc[0, "gr_school"], "École\t東京")
        self.assertEqual(reread.loc[2, "gr_pell_amt"], self.frame.loc[2, "gr_pell_amt"])
        self.assertEqual(reread.loc[2, "gr_pell_raw"], 'line1\nline2,"quoted"')

    def test_csv_reserved_null_token_cannot_silently_erase_literal_text(self) -> None:
        frame = self.frame.copy()
        frame.loc[0, "gr_school"] = "__FSA_NULL__"
        with self.assertRaises(ValueError):
            exports.write_csv(frame, self.root / "collision.csv.gz")

    def test_stata_roundtrip_checks_native_labels_long_strings_and_missingness(self) -> None:
        path = self.root / "data.dta"
        exports.write_stata(self.frame, path, self.metadata, self.labels)
        exports.verify_stata(path, self.frame, self.metadata, self.labels)
        with pd.io.stata.StataReader(path) as reader:
            native_labels = reader.variable_labels()
            native_values = reader.value_labels()
        self.assertEqual(native_labels["gr_opeid8"], "Research label: gr_opeid8")
        self.assertTrue(any(mapping == {0: "No", 1: "Yes"} for mapping in native_values.values()))
        raw = pd.read_stata(path, convert_categoricals=False)
        self.assertEqual(raw.loc[0, "gr_opeid8"], "00100000")
        self.assertEqual(raw.loc[0, "gr_school"], "École\t東京")
        self.assertEqual(len(raw.loc[0, "gr_pell_raw"]), 2300)
        self.assertTrue(pd.isna(raw.loc[2, "gr_strict"]))
        self.assertEqual(raw.loc[0, "gr_pell_amt"], 0)
        self.assertTrue(pd.isna(raw.loc[3, "gr_pell_amt"]))

    def test_stata_rejects_integer_that_would_round_when_converted_to_double(self) -> None:
        frame = self.frame.copy()
        frame.loc[0, "gr_pell_n"] = 2**53 + 1
        with self.assertRaises(ValueError):
            exports.write_stata(frame, self.root / "unsafe.dta", self.metadata, self.labels)

    def test_excel_treats_identifiers_and_formula_like_text_as_strings(self) -> None:
        path = self.root / "data.xlsx"
        frame = self.frame.copy()
        frame.loc[2, "gr_pell_amt"] = 123456789.12
        exports.write_excel(frame, path, self.metadata, self.labels, {"title": "Adversarial test fixture"})
        book = load_workbook(path, read_only=False, data_only=False)
        self.addCleanup(book.close)
        data = next(sheet for sheet in book if sheet.cell(1, 1).value == "unitid")
        headers = {cell.value: cell.column for cell in data[1]}
        self.assertEqual(data.cell(2, headers["gr_opeid8"]).value, "00100000")
        self.assertEqual(data.cell(2, headers["gr_opeid8"]).data_type, "s")
        self.assertEqual(data.cell(3, headers["gr_school"]).value, "=2+2")
        self.assertEqual(data.cell(3, headers["gr_school"]).data_type, "s")
        self.assertEqual(data.cell(3, headers["gr_pell_raw"]).value, "@SUM(A1)")
        self.assertEqual(data.cell(4, headers["gr_school"]).value, "NA")
        self.assertIsNone(data.cell(4, headers["gr_strict"]).value)
        self.assertEqual(data.cell(2, headers["gr_pell_amt"]).value, 0)
        self.assertIsNone(data.cell(5, headers["gr_pell_amt"]).value)
        self.assertIn("Codebook", book.sheetnames)
        self.assertIn("ValueLabels", book.sheetnames)
        self.assertIn("README", book.sheetnames)

    def test_excel_rejects_more_precision_than_native_excel_preserves(self) -> None:
        with self.assertRaises(ValueError):
            exports.write_excel(self.frame, self.root / "unsafe.xlsx", self.metadata, self.labels, {})

    def test_excel_preserves_currency_cents_and_counts_binary_roundoff(self) -> None:
        frame = self.frame.copy()
        frame.loc[2, "gr_pell_amt"] = 26343591.770000003
        path = self.root / "money_roundoff.xlsx"
        exports.write_excel(frame, path, self.metadata, self.labels, {})
        actual = exports.read_excel_data(path, list(frame))
        self.assertEqual(actual.loc[2, "gr_pell_amt"], 26343591.77)
        result = exports.verify_frame(frame, actual, self.metadata, "excel")
        self.assertEqual(result["excel_binary_rounding_cells_cents_preserved"], 1)
        with self.assertRaises(ValueError):
            exports.verify_frame(frame, actual, self.metadata, "csv")

    def test_excel_money_tolerance_does_not_authorize_noncent_precision_loss(self) -> None:
        frame = self.frame.copy()
        frame.loc[2, "gr_pell_amt"] = 0.12345678912345678
        with self.assertRaises(ValueError):
            exports.write_excel(frame, self.root / "unsafe_subcent.xlsx", self.metadata, self.labels, {})
        rounded = frame.copy()
        rounded.loc[2, "gr_pell_amt"] = float(format(frame.loc[2, "gr_pell_amt"], ".15g"))
        with self.assertRaises(ValueError):
            exports.verify_frame(frame, rounded, self.metadata, "excel")

    def test_excel_name_similarity_uses_bounded_native_precision(self) -> None:
        frame = self.frame.copy()
        frame.loc[2, "gr_pell_amt"] = 123456789.12
        frame["gr_name_similarity"] = pd.Series([0.7222222222222222, 0.0, 1.0, pd.NA, 0.5], dtype="Float64")
        metadata = pd.concat([self.metadata, pd.DataFrame([{
            "canonical_name": "grant__ipeds_fsa_hd_name_similarity",
            "portable_name": "gr_name_similarity",
            "variable_label": "Grant: IPEDS FSA name similarity",
            "description": "Name comparison diagnostic score, not a certified identity match.",
            "units": "ratio", "role": "quality_status", "source_family": "grant",
            "original_dtype": "Float64", "export_storage": "numeric", "category_kind": "",
        }])], ignore_index=True)
        path = self.root / "similarity_roundoff.xlsx"
        exports.write_excel(frame, path, metadata, self.labels, {})
        actual = exports.read_excel_data(path, list(frame))
        self.assertEqual(actual.loc[0, "gr_name_similarity"], 0.722222222222222)
        result = exports.verify_frame(frame, actual, metadata, "excel")
        self.assertTrue(result["all_values_passed"])
        corrupt = actual.copy()
        corrupt.loc[0, "gr_name_similarity"] = 0.72
        with self.assertRaises(ValueError):
            exports.verify_frame(frame, corrupt, metadata, "excel")

    def test_excel_rejects_unsupported_control_characters_without_silent_loss(self) -> None:
        frame = self.frame.copy()
        frame.loc[0, "gr_school"] = "Null\x00character"
        with self.assertRaises(ValueError):
            exports.write_excel(frame, self.root / "unsafe.xlsx", self.metadata, self.labels, {})

    def test_excel_rejects_text_beyond_native_cell_limit(self) -> None:
        frame = self.frame.copy()
        frame.loc[0, "gr_school"] = "x" * 32768
        with self.assertRaises(ValueError):
            exports.write_excel(frame, self.root / "unsafe.xlsx", self.metadata, self.labels, {})

    def test_parquet_carries_dataset_and_variable_metadata(self) -> None:
        path = self.root / "data.parquet"
        dataset_metadata = {"source_sha256": "a" * 64, "title": "Adversarial fixture"}
        exports.write_labeled_parquet(self.frame, path, self.metadata, self.labels, dataset_metadata)
        table = pq.read_table(path)
        exports.verify_frame(self.frame, table.to_pandas(), self.metadata, "parquet")
        serialized = str(table.schema.metadata) + str(table.schema.field("gr_opeid8").metadata)
        self.assertIn("source_sha256", serialized)
        self.assertIn("a" * 64, serialized)
        self.assertIn("Research label: gr_opeid8", serialized)
        status_meta = str(table.schema.field("gr_pell_status").metadata)
        self.assertIn("observed", status_meta)
        self.assertIn("suppressed", status_meta)

    def test_verifier_detects_value_and_missingness_corruption(self) -> None:
        for name, value in [("gr_opeid8", "100000"), ("gr_pell_amt", 999), ("gr_strict", 0)]:
            corrupt = self.frame.copy()
            row = 2 if name == "gr_strict" else 0
            corrupt.loc[row, name] = value
            with self.subTest(column=name), self.assertRaises(ValueError):
                exports.verify_frame(self.frame, corrupt, self.metadata, "csv")

    def _source_guard_fixture(self) -> tuple[Path, Path, Path]:
        release = self.root / "source"
        data_dir = release / "Panels/ipeds/unitid_research"
        data_dir.mkdir(parents=True)
        source = data_dir / "fsa_unitid_award_year_panel.parquet"
        source.write_bytes(b"Source fixture bytes; guard must run before parquet parsing")
        dictionary = data_dir / "dictionary.csv"
        dictionary.write_text("panel_column,definition\nunitid,IPEDS institution ID\n")
        (data_dir / "manifest.json").write_text(json.dumps({
            "all_conservation_passed": True,
            "output_sha256": {"dictionary": exports.sha256(dictionary)},
        }))
        (release / "build").mkdir()
        manifest = release / "build/research_release_manifest.json"
        manifest.write_text(json.dumps({
            "acceptance_passed": True,
            "release_version": "adversarial_test_fixture",
            "artifacts": [{"path": str(source.relative_to(release)), "sha256": exports.sha256(source)}],
        }))
        return release, source, manifest

    def test_export_rejects_unaccepted_release_before_reading_data(self) -> None:
        release, _, manifest = self._source_guard_fixture()
        content = json.loads(manifest.read_text())
        content["acceptance_passed"] = False
        manifest.write_text(json.dumps(content))
        with self.assertRaisesRegex(ValueError, "acceptance"):
            exports.export_release(release, self.root / "export")
        self.assertFalse((self.root / "export").exists())

    def test_export_rejects_source_checksum_drift_before_reading_data(self) -> None:
        release, source, _ = self._source_guard_fixture()
        source.write_bytes(source.read_bytes() + b" changed")
        with self.assertRaisesRegex(ValueError, "Source panel differs"):
            exports.export_release(release, self.root / "export")
        self.assertFalse((self.root / "export").exists())

    def test_export_rejects_dictionary_checksum_drift_before_reading_data(self) -> None:
        release, source, _ = self._source_guard_fixture()
        dictionary = source.parent / "dictionary.csv"
        dictionary.write_text(dictionary.read_text() + "other,changed\n")
        with self.assertRaisesRegex(ValueError, "dictionary or conservation"):
            exports.export_release(release, self.root / "export")
        self.assertFalse((self.root / "export").exists())


if __name__ == "__main__":
    unittest.main()
